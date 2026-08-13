from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.exports.data import (
    HR_STATUS_LABELS,
    RECOMMENDATION_LABELS,
    SOURCE_LABELS,
    ExportData,
    translate,
)

HEADER_FILL = PatternFill("solid", fgColor="1F2933")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F2933")
MUTED_FONT = Font(size=9, color="6B7280")
THIN = Side(style="thin", color="D7DCE2")
BORDER = Border(bottom=THIN)


def _sheet_names(language: str) -> tuple[str, str, str]:
    if language == "fr":
        return "Classement", "Critères", "Synthèse"
    return "Ranking", "Criteria", "Summary"


def _style_header(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _autosize(sheet, widths: dict[int, int]) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width


def _ranking_sheet(workbook: Workbook, data: ExportData, name: str) -> None:
    sheet = workbook.active
    sheet.title = name

    headers = [
        data.label("rank"),
        data.label("candidate"),
        data.label("email"),
        data.label("phone"),
        data.label("score"),
        data.label("ai_score"),
        data.label("manual_score"),
        data.label("recommendation"),
        data.label("hr_status"),
        data.label("source"),
        data.label("experience"),
        data.label("missing_must_haves"),
        *[f"{c.name} ({c.weight})" for c in data.criteria],
        data.label("notes"),
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))

    for candidate in data.candidates:
        sheet.append([
            candidate.rank,
            candidate.full_name,
            candidate.email or "",
            candidate.phone or "",
            candidate.score,
            candidate.ai_score,
            candidate.manual_score,
            translate(RECOMMENDATION_LABELS, data.language, candidate.recommendation),
            translate(HR_STATUS_LABELS, data.language, candidate.hr_status),
            translate(SOURCE_LABELS, data.language, candidate.source),
            candidate.years_experience,
            ", ".join(candidate.missing_must_haves),
            *[candidate.scores_by_criterion.get(c.id) for c in data.criteria],
            candidate.hr_notes or "",
        ])

    last_row = sheet.max_row
    first_criterion = 13
    last_criterion = first_criterion + len(data.criteria) - 1

    if last_row > 1:
        # Colour scale across every score column, so a weak criterion is visible
        # at a glance next to a strong overall score.
        score_range = f"E2:G{last_row}"
        criteria_range = (
            f"{get_column_letter(first_criterion)}2:"
            f"{get_column_letter(last_criterion)}{last_row}"
        )
        rule = ColorScaleRule(
            start_type="num", start_value=0, start_color="F4C7C3",
            mid_type="num", mid_value=50, mid_color="FDE9A9",
            end_type="num", end_value=100, end_color="C6E7CE",
        )
        sheet.conditional_formatting.add(score_range, rule)
        if data.criteria:
            sheet.conditional_formatting.add(criteria_range, rule)

        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"

    sheet.freeze_panes = "C2"
    _autosize(sheet, {1: 6, 2: 28, 3: 28, 4: 16, 5: 9, 6: 9, 7: 11, 8: 16, 9: 15, 10: 16, 11: 12, 12: 30})
    for column in range(first_criterion, last_criterion + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 14
    sheet.column_dimensions[get_column_letter(len(headers))].width = 40


def _criteria_sheet(workbook: Workbook, data: ExportData, name: str) -> None:
    sheet = workbook.create_sheet(name)
    headers = [
        data.label("criterion"),
        data.label("weight"),
        data.label("must_have"),
        data.label("summary"),
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))

    for criterion in data.criteria:
        sheet.append([
            criterion.name,
            criterion.weight,
            data.label("yes") if criterion.is_must_have else data.label("no"),
            criterion.description or "",
        ])
    for row in sheet.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row:
            cell.border = BORDER

    _autosize(sheet, {1: 34, 2: 10, 3: 14, 4: 70})


def _summary_sheet(workbook: Workbook, data: ExportData, name: str) -> None:
    sheet = workbook.create_sheet(name)
    scores = [c.score for c in data.candidates if c.score is not None]

    sheet["A1"] = f"{data.label('title')} — {data.session.title}"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = data.label("disclaimer")
    sheet["A2"].font = MUTED_FONT
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A2:D2")
    sheet.row_dimensions[2].height = 40

    rows: list[tuple[str, object]] = [
        (data.label("position"), data.session.position),
        (data.label("department"), data.session.department or "—"),
        (data.label("generated"), data.generated_at.strftime("%d/%m/%Y %H:%M UTC")),
        (data.label("scope"), data.label(f"scope_{data.scope}")),
        (data.label("candidates"), len(data.candidates)),
        (data.label("threshold"), data.session.score_threshold),
        (
            data.label("average"),
            round(sum(scores) / len(scores), 1) if scores else "—",
        ),
    ]

    start = 4
    sheet.cell(row=start, column=1, value=data.label("metric"))
    sheet.cell(row=start, column=2, value=data.label("value"))
    _style_header(sheet, start, 2)
    for offset, (key, value) in enumerate(rows, start=start + 1):
        sheet.cell(row=offset, column=1, value=key).font = Font(bold=True, size=10)
        sheet.cell(row=offset, column=2, value=value)

    breakdown_start = start + len(rows) + 2
    sheet.cell(row=breakdown_start, column=1, value=data.label("recommendation"))
    sheet.cell(row=breakdown_start, column=2, value=data.label("candidates"))
    _style_header(sheet, breakdown_start, 2)

    counts: dict[str, int] = {}
    for candidate in data.candidates:
        key = candidate.recommendation or "—"
        counts[key] = counts.get(key, 0) + 1
    for offset, (key, count) in enumerate(counts.items(), start=breakdown_start + 1):
        sheet.cell(
            row=offset, column=1, value=translate(RECOMMENDATION_LABELS, data.language, key)
        )
        sheet.cell(row=offset, column=2, value=count)

    distribution_start = breakdown_start + len(counts) + 2
    sheet.cell(row=distribution_start, column=1, value=data.label("distribution"))
    sheet.cell(row=distribution_start, column=2, value=data.label("candidates"))
    _style_header(sheet, distribution_start, 2)
    for offset, (low, high) in enumerate(
        ((0, 19), (20, 39), (40, 59), (60, 79), (80, 100)), start=distribution_start + 1
    ):
        sheet.cell(row=offset, column=1, value=f"{low}–{high}")
        sheet.cell(row=offset, column=2, value=sum(1 for s in scores if low <= s <= high))

    _autosize(sheet, {1: 34, 2: 30, 3: 20, 4: 20})


def render(data: ExportData) -> bytes:
    workbook = Workbook()
    ranking, criteria, summary = _sheet_names(data.language)

    _ranking_sheet(workbook, data, ranking)
    _criteria_sheet(workbook, data, criteria)
    _summary_sheet(workbook, data, summary)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
