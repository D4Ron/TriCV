from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.services.exports.data import (
    HR_STATUS_LABELS,
    RECOMMENDATION_LABELS,
    ExportData,
    translate,
)

INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ALERT = RGBColor(0xA8, 0x32, 0x32)


def _configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = INK
        style.font.bold = True


def _muted(document: Document, text: str, size: float = 9) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = MUTED


def _fact_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = str(value)
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    document.add_paragraph()


def render(data: ExportData) -> bytes:
    document = Document()
    _configure_styles(document)

    document.add_heading(data.label("title"), level=1)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(data.session.title)
    run.font.size = Pt(13)
    run.font.color.rgb = MUTED

    scores = [c.score for c in data.candidates if c.score is not None]
    _fact_table(document, [
        (data.label("position"), data.session.position),
        (data.label("department"), data.session.department or "—"),
        (data.label("generated"), data.generated_at.strftime("%d/%m/%Y %H:%M UTC")),
        (data.label("scope"), data.label(f"scope_{data.scope}")),
        (data.label("candidates"), f"{len(data.candidates)} / {data.total_in_session}"),
        (data.label("threshold"), f"{data.session.score_threshold}/100"),
        (
            data.label("average"),
            f"{sum(scores) / len(scores):.1f}/100" if scores else "—",
        ),
    ])

    if data.session.description:
        document.add_paragraph(data.session.description)

    _muted(document, data.label("disclaimer"), size=8.5)

    # --- criteria ----------------------------------------------------------
    document.add_heading(data.label("criteria"), level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for index, header in enumerate(
        (data.label("criterion"), data.label("weight"), data.label("must_have"),
         data.label("summary"))
    ):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    for criterion in data.criteria:
        cells = table.add_row().cells
        cells[0].text = criterion.name
        cells[1].text = f"{criterion.weight}/10"
        cells[2].text = data.label("yes") if criterion.is_must_have else data.label("no")
        cells[3].text = criterion.description or "—"

    # --- ranking -----------------------------------------------------------
    document.add_heading(data.label("ranking"), level=2)
    if not data.candidates:
        document.add_paragraph(data.label("no_candidates"))
    else:
        ranking = document.add_table(rows=1, cols=5)
        ranking.style = "Light Grid Accent 1"
        for index, header in enumerate(
            (data.label("rank"), data.label("candidate"), data.label("score"),
             data.label("recommendation"), data.label("hr_status"))
        ):
            cell = ranking.rows[0].cells[index]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True

        for candidate in data.candidates:
            cells = ranking.add_row().cells
            cells[0].text = str(candidate.rank)
            cells[1].text = candidate.full_name
            cells[2].text = "—" if candidate.score is None else f"{candidate.score:.0f}"
            cells[3].text = translate(
                RECOMMENDATION_LABELS, data.language, candidate.recommendation
            )
            cells[4].text = translate(HR_STATUS_LABELS, data.language, candidate.hr_status)

    # --- one section per candidate ----------------------------------------
    if data.candidates:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading(data.label("detail"), level=2)

        for candidate in data.candidates:
            document.add_heading(f"{candidate.rank}. {candidate.full_name}", level=3)

            _fact_table(document, [
                (
                    data.label("score"),
                    "—" if candidate.score is None else f"{candidate.score:.1f}/100",
                ),
                (
                    data.label("recommendation"),
                    translate(RECOMMENDATION_LABELS, data.language, candidate.recommendation),
                ),
                (
                    data.label("hr_status"),
                    translate(HR_STATUS_LABELS, data.language, candidate.hr_status),
                ),
                (data.label("email"), candidate.email or "—"),
                (data.label("phone"), candidate.phone or "—"),
                (
                    data.label("experience"),
                    f"{candidate.years_experience} {data.label('years')}"
                    if candidate.years_experience is not None
                    else "—",
                ),
                (data.label("education"), candidate.education_level or "—"),
            ])

            if candidate.summary:
                document.add_paragraph(candidate.summary)

            if candidate.missing_must_haves:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(
                    f"{data.label('missing_must_haves')}: "
                    + ", ".join(candidate.missing_must_haves)
                )
                run.font.bold = True
                run.font.color.rgb = ALERT

            for key, items in (("strengths", candidate.strengths), ("gaps", candidate.gaps)):
                if items:
                    document.add_paragraph(data.label(key), style="Heading 3")
                    for item in items:
                        document.add_paragraph(item, style="List Bullet")

            if candidate.scores_by_criterion:
                document.add_paragraph(data.label("criteria"), style="Heading 3")
                detail = document.add_table(rows=1, cols=3)
                detail.style = "Light List Accent 1"
                for index, header in enumerate(
                    (data.label("criterion"), data.label("score"), data.label("justification"))
                ):
                    cell = detail.rows[0].cells[index]
                    cell.text = header
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True

                for criterion in data.criteria:
                    score = candidate.scores_by_criterion.get(criterion.id)
                    cells = detail.add_row().cells
                    cells[0].text = criterion.name + (" *" if criterion.is_must_have else "")
                    cells[1].text = "—" if score is None else f"{score:.0f}"
                    cells[2].text = candidate.justifications.get(criterion.id) or "—"

            if candidate.hr_notes:
                document.add_paragraph(data.label("notes"), style="Heading 3")
                document.add_paragraph(candidate.hr_notes)

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Kapi Consult · TriCV"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
