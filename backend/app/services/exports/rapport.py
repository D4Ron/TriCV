"""Le rapport de recrutement, rendu dans les formats que les clients demandent.

Quatre sorties depuis un seul contenu : DOCX pour la remise habituelle et pour
que le client puisse reprendre le texte, PDF pour la version qui fait foi, ODT
pour les administrations qui travaillent sous LibreOffice, TXT pour ceux qui
imposent un dépôt en texte brut.

Les sections arrivent déjà rédigées et relues. Ce module ne décide de rien ; il
met en page. C'est délibéré : le fond et la forme se sont séparés au moment où
un client a demandé sa propre trame, et les mélanger de nouveau obligerait à
réécrire la génération pour chaque format imposé.

Le document remis ne se limite pas au corps. Il s'ouvre sur une page de garde à
l'en-tête du cabinet, puis sur un sommaire, et ses titres sont numérotés — voir
`frontispice`, qui tient ce qui est commun aux quatre formats. Chacun le dessine
ici avec ses propres moyens : un champ TOC pour Word, un `text:table-of-content`
pour LibreOffice, une table des matières calculée en deux passes pour le PDF,
une liste simple pour le texte brut.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

from app.services.exports import frontispice
from app.services.exports.frontispice import Couverture

# Charte Kapi Consult : bleu de la marque, or des filets, gris du texte
# secondaire. Les mêmes valeurs que l'interface, pour qu'un rapport imprimé et
# un écran ne se contredisent pas.
BLEU = RGBColor(0x1E, 0x22, 0x99)
OR = RGBColor(0xB8, 0x89, 0x2A)
GRIS = RGBColor(0x6B, 0x70, 0x80)

BLEU_HEX = "#1E2299"
OR_HEX = "#B8892A"
GRIS_HEX = "#6B7080"

# Fonds des lignes remarquables. Assez pâles pour rester lisibles à
# l'impression en noir et blanc, où seule la graisse distinguera les lignes.
FOND_ENTETE = "E8E9F5"
FOND_RUBRIQUE = "F2F2F6"
FOND_TOTAL = "EDEEF8"

GENRE_RUBRIQUE = "rubrique"
GENRE_DETAIL = "detail"
GENRE_TOTAL = "total"

TITRE_SOMMAIRE = "Sommaire"


# --- lecture des tableaux structurés ----------------------------------------
#
# Une section calculée porte `tableaux` : colonnes, lignes, et le genre de
# chaque ligne. Un rapport produit avant que cette structure n'existe ne les a
# pas — il retombe alors sur `contenu`, son texte aligné à l'espace. C'est le
# seul cas où les colonnes se décalent, et il ne concerne que les rapports
# déjà en base.


def _tableaux_de(section: dict) -> list[dict]:
    tableaux = section.get("tableaux")
    return [t for t in tableaux if isinstance(t, dict)] if isinstance(tableaux, list) else []


def _colonnes(tableau: dict) -> list[dict]:
    return [c for c in (tableau.get("colonnes") or ()) if isinstance(c, dict)]


def _lignes(tableau: dict) -> list[dict]:
    return [l for l in (tableau.get("lignes") or ()) if isinstance(l, dict)]


def _cellules(ligne: dict, largeur: int) -> list[str]:
    valeurs = [str(c) for c in (ligne.get("cellules") or ())]
    # Une ligne plus courte que l'en-tête ne doit pas faire tomber l'export.
    return (valeurs + [""] * largeur)[:largeur]


def _niveau(section: dict) -> int:
    """1 ou 2. Le rapport remis est hiérarchisé, pas une liste de titres."""
    try:
        niveau = int(section.get("niveau") or 1)
    except (TypeError, ValueError):
        return 1
    return 2 if niveau >= 2 else 1


def _a_rendre(section: dict) -> bool:
    """Une section vide ne paraît pas — sauf si elle porte des sous-sections.

    « Méthodologie » n'a pas de texte à elle. La sauter laisserait
    « Présélection » et « Critères éliminatoires » sans rattachement, au même
    rang que tout le reste.
    """
    return bool(
        str(section.get("contenu") or "").strip()
        or str(section.get("contenu_apres") or "").strip()
        or _tableaux_de(section)
        or section.get("porteur")
    )


def preparer(
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None,
) -> tuple[list[dict], list[frontispice.Entree], Couverture]:
    """Les sections effectivement rendues, leurs numéros, et la page de garde.

    Le filtrage précède la numérotation : une section laissée vide ne paraît
    pas, elle ne doit donc pas consommer un numéro et laisser un trou à la
    place du II.
    """
    rendues = [s for s in sections if _a_rendre(s)]
    entrees = frontispice.numeroter(rendues)
    garde = couverture or frontispice.couverture_par_defaut(titre, sous_titre, genere_le)
    return rendues, entrees, garde


# Les marques dont la rédaction préfixe un élément d'énumération. Le modèle
# écrit « - Formation académique (7 points) : … » ; le document remis, lui,
# porte une vraie liste à puces, avec le retrait et l'alignement que Word et
# LibreOffice savent tenir. Rendre le tiret tel quel donnait une suite de
# paragraphes ordinaires commençant par un signe moins, que le lecteur
# reconstitue en liste de lui-même — et qui se replient sans retrait dès
# qu'un élément dépasse la ligne.
_PUCES = ("- ", "– ", "— ", "• ", "* ")


def paragraphes_de(section: dict, cle: str = "contenu") -> list[tuple[str, bool]]:
    """La prose d'une section, découpée en (texte, est-une-puce).

    Partagée par les quatre formats : un même contenu doit produire la même
    liste dans le DOCX, le PDF, l'ODT et l'aperçu. La reconnaissance vit ici
    pour qu'aucun format ne puisse en diverger.

    `cle` vaut « contenu » — la prose qui précède les tableaux — ou
    « contenu_apres », le commentaire des chiffres, qui se lit après eux.
    """
    sortie: list[tuple[str, bool]] = []
    for ligne in str(section.get(cle) or "").split("\n"):
        texte = ligne.strip()
        if not texte:
            continue
        for puce in _PUCES:
            if texte.startswith(puce):
                sortie.append((texte[len(puce):].strip(), True))
                break
        else:
            sortie.append((texte, False))
    return sortie


# La mention qui clôt le rapport du cabinet, sous le tableau des résultats.
NOTE_ANNEXE = (
    "NB : Le détail des notes obtenues par chaque candidat est annexé au "
    "présent rapport."
)


def _note_de_fin(section: dict) -> str:
    """La mention d'annexe, sous le dernier tableau de résultats."""
    return NOTE_ANNEXE if section.get("code") == "RESULTATS_ENTRETIENS" else ""


# --- DOCX -------------------------------------------------------------------


def _ombrer(cellule, fond: str) -> None:
    """Un fond de cellule. python-docx ne l'expose pas : on pose le XML."""
    proprietes = cellule._tc.get_or_add_tcPr()
    ombre = OxmlElement("w:shd")
    ombre.set(qn("w:val"), "clear")
    ombre.set(qn("w:fill"), fond)
    proprietes.append(ombre)


def _ecrire_cellule(
    cellule, texte: str, *, gras: bool, droite: bool, retrait: bool
) -> None:
    paragraphe = cellule.paragraphs[0]
    paragraphe.paragraph_format.space_after = Pt(2)
    paragraphe.paragraph_format.space_before = Pt(2)
    if droite:
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if retrait:
        paragraphe.paragraph_format.left_indent = Pt(12)
    run = paragraphe.add_run(texte)
    run.font.size = Pt(9.5)
    run.font.bold = gras


def _ajuster_a_la_page(table) -> None:
    """« Ajuster à la fenêtre » plutôt qu'« ajuster au contenu ».

    python-docx pose `tblW type="auto"`, que Word lit comme « ajuster au
    contenu » : la table s'élargit autant qu'il le faut. Le tableau des
    préqualifiés porte huit colonnes — nom, âge, diplôme, pays, note, rang,
    téléphone, courriel — et débordait donc de la colonne de texte, une moitié
    hors de la page imprimée.

    Une largeur de 100 % le ramène dans les marges et laisse Word replier le
    contenu des cellules, ce qui est le comportement voulu : un courriel long
    passe à la ligne, il ne pousse pas la page.
    """
    proprietes = table._tbl.tblPr
    for ancien in proprietes.findall(qn("w:tblW")):
        proprietes.remove(ancien)
    largeur = OxmlElement("w:tblW")
    largeur.set(qn("w:type"), "pct")
    largeur.set(qn("w:w"), "5000")  # 5000 cinquantièmes de pour cent = 100 %
    # `w:tblW` suit `w:tblStyle` et précède `w:jc` dans la séquence du schéma.
    proprietes.insert_element_before(
        largeur,
        "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
        "w:tblDescription", "w:tblPrChange",
    )


def _table_docx(document, tableau: dict) -> None:
    colonnes = _colonnes(tableau)
    lignes = _lignes(tableau)
    if not colonnes or not lignes:
        return

    table = document.add_table(rows=1, cols=len(colonnes))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # `autofit` laisse Word répartir les colonnes sur le contenu réel, ce qui
    # évite d'imposer des largeurs qui coupent un intitulé long.
    table.autofit = True
    _ajuster_a_la_page(table)

    for index, colonne in enumerate(table.rows[0].cells):
        _ombrer(colonne, FOND_ENTETE)
        _ecrire_cellule(
            colonne,
            str(colonnes[index].get("libelle") or ""),
            gras=True,
            droite=bool(colonnes[index].get("numerique")),
            retrait=False,
        )

    for ligne in lignes:
        genre = str(ligne.get("genre") or "")
        cellules = table.add_row().cells
        for index, valeur in enumerate(_cellules(ligne, len(colonnes))):
            if genre == GENRE_RUBRIQUE:
                _ombrer(cellules[index], FOND_RUBRIQUE)
            elif genre == GENRE_TOTAL:
                _ombrer(cellules[index], FOND_TOTAL)
            _ecrire_cellule(
                cellules[index],
                valeur,
                gras=genre in (GENRE_RUBRIQUE, GENRE_TOTAL),
                droite=bool(colonnes[index].get("numerique")),
                # Le sous-critère se lit sous sa rubrique, pas à côté d'elle.
                retrait=genre == GENRE_DETAIL and index == 0,
            )


def _section_docx(document, section: dict) -> None:
    """Le corps d'une section : sa prose, puis ses tables.

    Dans cet ordre, parce que c'est celui du document remis : une phrase
    annonce le tableau, le tableau suit.
    """
    for paragraphe, puce in paragraphes_de(section):
        document.add_paragraph(paragraphe, style="List Bullet" if puce else None)

    for tableau in _tableaux_de(section):
        legende = str(tableau.get("titre") or "").strip()
        if legende:
            for ligne in legende.split("\n"):
                paragraphe = document.add_paragraph()
                paragraphe.paragraph_format.space_after = Pt(3)
                run = paragraphe.add_run(ligne)
                run.font.bold = True
                run.font.size = Pt(10)
        _table_docx(document, tableau)
        document.add_paragraph()

    # Le commentaire des chiffres, après les chiffres.
    for paragraphe, puce in paragraphes_de(section, "contenu_apres"):
        document.add_paragraph(paragraphe, style="List Bullet" if puce else None)

    note = _note_de_fin(section)
    if note:
        paragraphe = document.add_paragraph()
        run = paragraphe.add_run(note)
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = GRIS


# --- DOCX : page de garde, sommaire, pied de page ---------------------------
#
# Word ne sait pas numéroter une page ni dresser une table des matières depuis
# un contenu figé : il le fait au moment où le document s'ouvre, à partir de
# *champs*. Un champ est une suite de runs encadrée par deux marques —
# `begin`…`separate`…`end` — dont python-docx n'expose rien. On pose le XML.
#
# Entre `separate` et `end` vit le *résultat mis en cache* : ce qu'affichent
# les lecteurs qui n'actualisent pas. On y écrit les titres numérotés, sans
# numéro de page. Un sommaire sans pagination reste un sommaire ; un champ vide
# ressemble à un document abîmé.


def _marque_de_champ(genre: str, sale: bool = False):
    marque = OxmlElement("w:fldChar")
    marque.set(qn("w:fldCharType"), genre)
    if sale:
        # « dirty » demande au lecteur d'actualiser ce champ à l'ouverture.
        marque.set(qn("w:dirty"), "true")
    return marque


def _instruction(texte: str):
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = texte
    return instruction


def _champ_simple(paragraphe, instruction: str, cache: str = "") -> None:
    """Un champ qui tient dans un paragraphe — un numéro de page, par exemple."""
    debut = paragraphe.add_run()
    debut._r.append(_marque_de_champ("begin"))
    milieu = paragraphe.add_run()
    milieu._r.append(_instruction(instruction))
    separation = paragraphe.add_run()
    separation._r.append(_marque_de_champ("separate"))
    paragraphe.add_run(cache)
    fin = paragraphe.add_run()
    fin._r.append(_marque_de_champ("end"))


# Les éléments de `word/settings.xml` qui, dans le schéma OOXML, viennent après
# `w:updateFields`. Même remarque que pour `w:pBdr` plus haut : Word tolère
# l'ordre inverse, on range par conformité et non pour réparer une panne.
# Ajouter à la fin — le réflexe — posait l'élément après tous ceux-ci.
_APRES_UPDATE_FIELDS = (
    "w:compat",
    "w:docVars",
    "w:rsids",
    "w:mathPr",
    "w:themeFontLang",
    "w:clrSchemeMapping",
    "w:doNotAutoCompressPictures",
    "w:shapeDefaults",
    "w:decimalSymbol",
    "w:listSeparator",
)


def _actualiser_les_champs(document) -> None:
    """Demande au lecteur d'actualiser la table des matières à l'ouverture.

    Sans cela, Word affiche le contenu mis en cache jusqu'à ce que quelqu'un
    pense à faire un clic droit dessus — et le rapport part au client avec un
    sommaire sans pages.
    """
    reglages = document.settings.element
    drapeau = OxmlElement("w:updateFields")
    drapeau.set(qn("w:val"), "true")

    for nom in _APRES_UPDATE_FIELDS:
        suivant = reglages.find(qn(nom))
        if suivant is not None:
            suivant.addprevious(drapeau)
            return
    reglages.append(drapeau)


# Ce qui, dans le schéma OOXML, suit `w:pBdr` à l'intérieur de `w:pPr`.
#
# `w:pPr` n'est pas un sac d'attributs : le schéma en décrit la *séquence*.
# python-docx range de lui-même ce qu'il sait poser ; `w:pBdr`, qu'il ne
# connaît pas, se retrouvait ajouté à la fin — donc après le `w:spacing` et le
# `w:jc` posés juste avant.
#
# Mesuré avant de conclure : Word 16 ouvre le document **sans broncher** dans
# les deux ordres, même avec la réparation automatique désactivée. Ce n'était
# donc pas la cause du rapport inexploitable, contrairement à ce qu'on a cru.
# On range quand même, parce que le document est conforme ou ne l'est pas, et
# que rien ne garantit la même indulgence chez un validateur, une bibliothèque
# tierce ou une version future.
_APRES_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _filet(paragraphe, couleur: str, epaisseur: int = 12) -> None:
    """Le filet or sous la marque. Une bordure basse de paragraphe."""
    proprietes = paragraphe._p.get_or_add_pPr()
    bordures = OxmlElement("w:pBdr")
    bas = OxmlElement("w:bottom")
    bas.set(qn("w:val"), "single")
    bas.set(qn("w:sz"), str(epaisseur))
    bas.set(qn("w:space"), "4")
    bas.set(qn("w:color"), couleur)
    bordures.append(bas)
    proprietes.insert_element_before(bordures, *_APRES_PBDR)


def _ligne(document, texte: str, *, taille: float, gras=False, couleur=None,
           centre=True, avant=0, apres=4, majuscules=False):
    paragraphe = document.add_paragraph()
    if centre:
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraphe.paragraph_format.space_before = Pt(avant)
    paragraphe.paragraph_format.space_after = Pt(apres)
    if texte:
        run = paragraphe.add_run(texte.upper() if majuscules else texte)
        run.font.size = Pt(taille)
        run.font.bold = gras
        if couleur is not None:
            run.font.color.rgb = couleur
    return paragraphe


def _fond_de_garde_docx(document) -> None:
    """Pose le papier à en-tête en fond de la première page.

    Word n'a pas de « fond de page » par section. L'usage, et ce que fait le
    document du cabinet, est d'ancrer une image en pleine page ; le plus court
    chemin depuis python-docx est de la poser dans le **pied de page de
    première page**, dont la distance au bord est ramenée à zéro. Elle se
    dessine alors sous le texte, d'un bord à l'autre.
    """
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.footer_distance = 0
    section.header_distance = 0

    pied = section.first_page_footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pied.paragraph_format.space_before = Pt(0)
    pied.paragraph_format.space_after = Pt(0)
    run = pied.add_run()
    run.add_picture(
        io.BytesIO(frontispice.entete_jpg()),
        width=section.page_width,
        height=section.page_height,
    )
    # Remonter l'image du bas de page jusqu'au haut de la feuille : sans ce
    # décalage, elle commence là où commence le pied, c'est-à-dire en bas.
    _remonter_en_pleine_page(run, section)


def _remonter_en_pleine_page(run, section) -> None:
    """Transforme l'image en objet flottant calé sur le coin de la page.

    python-docx ne pose que des images « en ligne ». Une image en ligne dans
    un pied de page reste dans le pied ; il faut la convertir en ancrage
    absolu, relatif à la *page*, pour qu'elle couvre la feuille entière.
    """
    dessin = run._r.find(qn("w:drawing"))
    if dessin is None:
        return
    en_ligne = dessin.find(qn("wp:inline"))
    if en_ligne is None:
        return

    ancre = OxmlElement("wp:anchor")
    for cle, valeur in (
        ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
        ("simplePos", "0"), ("relativeHeight", "0"), ("behindDoc", "1"),
        ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1"),
    ):
        ancre.set(cle, valeur)

    pos_simple = OxmlElement("wp:simplePos")
    pos_simple.set("x", "0")
    pos_simple.set("y", "0")
    ancre.append(pos_simple)

    for nom, sens in (("wp:positionH", "page"), ("wp:positionV", "page")):
        pos = OxmlElement(nom)
        pos.set("relativeFrom", sens)
        decalage = OxmlElement("wp:posOffset")
        decalage.text = "0"
        pos.append(decalage)
        ancre.append(pos)

    # `extent`, `docPr` et le graphique se reprennent tels quels.
    for enfant in list(en_ligne):
        ancre.append(enfant)
    enveloppe = OxmlElement("wp:wrapNone")
    # `wrapNone` se place après `extent`/`effectExtent`, avant `docPr`.
    ancre.insert(list(ancre).index(ancre.find(qn("wp:docPr"))), enveloppe)

    dessin.remove(en_ligne)
    dessin.append(ancre)


def _page_de_garde_docx(document, garde: Couverture) -> None:
    """Ce qui se pose SUR le papier à en-tête : l'objet, le titre, le mois.

    Le fond porte déjà la marque, les filets, les coordonnées et la mention de
    confidentialité : les recomposer ici les ferait paraître deux fois.
    """
    _fond_de_garde_docx(document)

    # Descendre jusqu'au panneau clair de l'en-tête, au-dessus de la bande
    # d'images qui court à mi-hauteur.
    _ligne(document, "", taille=11, avant=0, apres=0)
    _ligne(document, "", taille=11, avant=120, apres=0)

    if garde.objet_affiche:
        _ligne(
            document, garde.objet_affiche, taille=13, gras=True, couleur=BLEU, apres=24
        )
    _ligne(document, garde.titre_document, taille=22, gras=True, couleur=BLEU, apres=10)
    if garde.client:
        _ligne(document, garde.client, taille=13, couleur=GRIS, apres=2)
    if garde.reference:
        _ligne(document, f"Réf. {garde.reference}", taille=10, couleur=GRIS, apres=2)
    if garde.mois:
        _ligne(document, garde.mois, taille=12, gras=True, couleur=GRIS, avant=30)


def _sommaire_docx(document, entrees: list[frontispice.Entree]) -> None:
    """Le sommaire : un champ TOC, avec les titres en cache.

    Le titre « Sommaire » est un paragraphe ordinaire, non un style de titre.
    Un titre en bonne et due forme se recenserait lui-même à la première ligne
    de sa propre table des matières.
    """
    _ligne(document, TITRE_SOMMAIRE, taille=16, gras=True, couleur=BLEU, apres=12)

    ouverture = document.add_paragraph()
    debut = ouverture.add_run()
    debut._r.append(_marque_de_champ("begin", sale=True))
    milieu = ouverture.add_run()
    milieu._r.append(_instruction(' TOC \\o "1-2" \\h \\z \\u '))
    separation = ouverture.add_run()
    separation._r.append(_marque_de_champ("separate"))

    for entree in entrees:
        cache = document.add_paragraph()
        cache.paragraph_format.space_after = Pt(3)
        if entree.niveau == 2:
            cache.paragraph_format.left_indent = Pt(18)
        run = cache.add_run(entree.titre_complet)
        run.font.size = Pt(10.5)
        run.font.bold = entree.niveau == 1

    fermeture = document.add_paragraph()
    fin = fermeture.add_run()
    fin._r.append(_marque_de_champ("end"))


def _pieds_de_page_docx(document, garde: Couverture) -> None:
    """Le bandeau du cabinet au pied des pages du corps, et la pagination.

    Le pied de **première page** n'est pas touché ici : il porte le papier à
    en-tête, posé par `_fond_de_garde_docx`.
    """
    section = document.sections[0]
    section.different_first_page_header_footer = True
    # De quoi loger le bandeau et la pagination sans que la dernière ligne du
    # texte ne vienne mordre dessus.
    section.bottom_margin = Cm(3.4)

    courant = section.footer.paragraphs[0]
    courant.alignment = WD_ALIGN_PARAGRAPH.CENTER
    courant.paragraph_format.space_after = Pt(0)
    numero = courant.add_run(f"{frontispice.MENTION_CONFIDENTIEL}   —   page ")
    numero.font.size = Pt(8)
    numero.font.color.rgb = GRIS
    _champ_simple(courant, " PAGE ")
    for run in courant.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS

    bandeau = section.footer.add_paragraph()
    bandeau.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bandeau.paragraph_format.space_before = Pt(2)
    bandeau.paragraph_format.space_after = Pt(0)
    largeur = section.page_width - section.left_margin - section.right_margin
    bandeau.add_run().add_picture(
        io.BytesIO(frontispice.bandeau_png()),
        width=largeur,
        height=int(largeur * frontispice.PROPORTION_BANDEAU),
    )


def rendre_docx(
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None = None,
) -> bytes:
    rendues, entrees, garde = preparer(
        titre, sous_titre, sections, genere_le, couverture
    )
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    # Le titre du rapport vit désormais sur la page de garde : « Heading 1 » est
    # libre pour les sections, ce qui est aussi ce que le champ TOC recense
    # (`\o "1-2"`). Une hiérarchie décalée d'un cran laissait la table des
    # matières vide.
    for nom, taille in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = document.styles[nom]
        style.font.name = "Calibri"
        style.font.size = Pt(taille)
        style.font.color.rgb = BLEU
        style.font.bold = True

    _pieds_de_page_docx(document, garde)
    _page_de_garde_docx(document, garde)

    # Une section Word neuve pour que la page de garde garde son pied à elle —
    # les coordonnées du cabinet — sans l'imposer au reste du document.
    # `add_section` recopie les réglages de la précédente, dont le pied de
    # première page : sans cette ligne, le sommaire reprenait l'adresse du
    # cabinet au lieu de sa pagination.
    suite = document.add_section(WD_SECTION.NEW_PAGE)
    suite.different_first_page_header_footer = False
    # La section neuve recopie les marges de la précédente, sauf si on les
    # repose : le bandeau du pied a besoin de la même place ici.
    suite.bottom_margin = Cm(3.4)
    _sommaire_docx(document, entrees)

    document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    for section, entree in zip(rendues, entrees):
        document.add_heading(entree.titre_complet, level=entree.niveau)
        _section_docx(document, section)

    _actualiser_les_champs(document)

    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


# --- PDF --------------------------------------------------------------------


# A4 moins les marges posées sur le SimpleDocTemplate plus bas.
MARGE = 20 * mm
LARGEUR_UTILE = A4[0] - 2 * MARGE


def _largeurs(tableau: dict, largeur_totale: float) -> list[float]:
    """Répartit la largeur sur les colonnes, au prorata de leur contenu.

    Bornée : une colonne de notes n'a pas besoin de plus que sa légende, et
    une colonne d'adresses ne doit pas écraser toutes les autres. Le reste se
    replie sur plusieurs lignes, ce que `Paragraph` sait faire dans une cellule.
    """
    colonnes = _colonnes(tableau)
    besoins = []
    for index, colonne in enumerate(colonnes):
        longueurs = [len(str(colonne.get("libelle") or ""))]
        longueurs += [
            len(_cellules(ligne, len(colonnes))[index]) for ligne in _lignes(tableau)
        ]
        besoins.append(max(6, min(46, max(longueurs))))

    total = sum(besoins) or 1
    return [largeur_totale * besoin / total for besoin in besoins]


def _table_pdf(tableau: dict, styles: dict, largeur_totale: float) -> list:
    colonnes = _colonnes(tableau)
    lignes = _lignes(tableau)
    legende = str(tableau.get("titre") or "").strip()

    elements: list = []
    for ligne in legende.split("\n") if legende else []:
        elements.append(Paragraph(escape(ligne), styles["legende"]))
    if not colonnes or not lignes:
        return elements

    def cellule(texte: str, *, style: str) -> Paragraph:
        return Paragraph(escape(texte), styles[style])

    donnees = [
        [
            cellule(
                str(c.get("libelle") or ""),
                style="entete_droite" if c.get("numerique") else "entete",
            )
            for c in colonnes
        ]
    ]
    decorations = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{FOND_ENTETE}")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CBD8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]

    for rang, ligne in enumerate(lignes, start=1):
        genre = str(ligne.get("genre") or "")
        gras = genre in (GENRE_RUBRIQUE, GENRE_TOTAL)
        valeurs = _cellules(ligne, len(colonnes))
        rendues = []
        for index, valeur in enumerate(valeurs):
            droite = bool(colonnes[index].get("numerique"))
            if genre == GENRE_DETAIL and index == 0:
                style = "cellule_retrait"
            elif gras:
                style = "cellule_grasse_droite" if droite else "cellule_grasse"
            else:
                style = "cellule_droite" if droite else "cellule"
            rendues.append(cellule(valeur, style=style))
        donnees.append(rendues)

        if genre == GENRE_RUBRIQUE:
            decorations.append(
                ("BACKGROUND", (0, rang), (-1, rang), colors.HexColor(f"#{FOND_RUBRIQUE}"))
            )
        elif genre == GENRE_TOTAL:
            decorations.append(
                ("BACKGROUND", (0, rang), (-1, rang), colors.HexColor(f"#{FOND_TOTAL}"))
            )

    table = Table(
        donnees,
        colWidths=_largeurs(tableau, largeur_totale),
        # L'en-tête se répète en haut de chaque page : un tableau de vingt
        # préqualifiés se lit sur deux pages, et la seconde sans en-tête
        # oblige à revenir en arrière pour savoir quelle colonne est laquelle.
        repeatRows=1,
    )
    table.setStyle(TableStyle(decorations))
    elements.append(table)
    elements.append(Spacer(1, 10))
    # Une légende ne doit pas rester seule en bas de page, séparée de sa table.
    return [KeepTogether(elements)] if len(donnees) <= 12 else elements


class _DocumentRapport(SimpleDocTemplate):
    """Un document qui signale ses titres à la table des matières.

    Le PDF est le seul des quatre formats où la pagination du sommaire est
    calculée par nous : Word et LibreOffice la refont à l'ouverture, le texte
    brut n'en a pas. Elle demande deux passes — une pour savoir sur quelle page
    tombe chaque titre, une pour écrire le sommaire avec — ce que
    `multiBuild` orchestre à condition qu'on lui dise où sont les titres.
    """

    _NIVEAUX = {"section": 0, "sous_section": 1}

    def afterFlowable(self, flowable) -> None:  # noqa: N802 — nom imposé
        if not isinstance(flowable, Paragraph):
            return
        niveau = self._NIVEAUX.get(flowable.style.name)
        if niveau is not None:
            self.notify("TOCEntry", (niveau, flowable.getPlainText(), self.page))


def _dessiner_pied(toile, texte_gauche: str, texte_droite: str) -> None:
    toile.saveState()
    toile.setFont("Helvetica", 7.5)
    toile.setFillColor(colors.HexColor(GRIS_HEX))
    if texte_gauche:
        toile.drawString(MARGE, 12 * mm, texte_gauche)
    if texte_droite:
        toile.drawRightString(A4[0] - MARGE, 12 * mm, texte_droite)
    toile.setStrokeColor(colors.HexColor("#DCDFEC"))
    toile.setLineWidth(0.4)
    toile.line(MARGE, 15 * mm, A4[0] - MARGE, 15 * mm)
    toile.restoreState()


def _pied_de_garde(garde: Couverture):
    """La première page : le papier à en-tête du cabinet, page pleine.

    Ce n'est pas une imitation. C'est l'image que le cabinet pose lui-même en
    fond de la première page de ses rapports — marque, filets, coordonnées et
    mention de confidentialité compris. On la dessine sous le texte, d'un bord
    à l'autre, et la garde n'a donc plus ni adresse ni filet à composer.
    """

    def dessiner(toile, _document) -> None:
        toile.saveState()
        if frontispice.pieces_presentes():
            toile.drawImage(
                ImageReader(io.BytesIO(frontispice.entete_jpg())),
                0,
                0,
                width=A4[0],
                height=A4[1],
                preserveAspectRatio=False,
                anchor="c",
            )
        else:
            # Dépôt sans les pièces : la garde reste lisible, composée en
            # texte, plutôt que de sortir blanche.
            toile.setStrokeColor(colors.HexColor(OR_HEX))
            toile.setLineWidth(1.2)
            toile.line(MARGE, 34 * mm, A4[0] - MARGE, 34 * mm)
            toile.setFont("Helvetica-Bold", 7.5)
            toile.setFillColor(colors.HexColor(GRIS_HEX))
            toile.drawString(MARGE, 29 * mm, frontispice.NOM_CABINET)
            toile.setFont("Helvetica", 7)
            for rang, ligne in enumerate(frontispice.ADRESSE):
                toile.drawString(MARGE, 25 * mm - rang * 3.6 * mm, ligne)
            toile.setFont("Helvetica-Bold", 8)
            toile.drawRightString(
                A4[0] - MARGE, 29 * mm, frontispice.MENTION_CONFIDENTIEL
            )
        toile.restoreState()

    return dessiner


def _pied_courant(garde: Couverture):
    """Le bandeau du cabinet au pied de chaque page, et le numéro.

    Le document remis ne met pas une ligne de texte en pied de page : il y met
    son bandeau « Nous développons vos métiers ». Le numéro de page se pose
    au-dessus, à droite, là où il ne le recouvre pas.
    """

    def dessiner(toile, _document) -> None:
        toile.saveState()
        if frontispice.pieces_presentes():
            largeur = A4[0] - 2 * MARGE
            hauteur = largeur * frontispice.PROPORTION_BANDEAU
            toile.drawImage(
                ImageReader(io.BytesIO(frontispice.bandeau_png())),
                MARGE,
                11 * mm,
                width=largeur,
                height=hauteur,
                preserveAspectRatio=True,
                anchor="sw",
                mask="auto",
            )
            toile.setFont("Helvetica", 7.5)
            toile.setFillColor(colors.HexColor(GRIS_HEX))
            toile.drawRightString(
                A4[0] - MARGE,
                11 * mm + hauteur + 2 * mm,
                f"{frontispice.MENTION_CONFIDENTIEL} — page {toile.getPageNumber()}",
            )
        else:
            _dessiner_pied(
                toile,
                f"{garde.titre_document} — {frontispice.MENTION_CONFIDENTIEL}",
                f"page {toile.getPageNumber()}",
            )
        toile.restoreState()

    return dessiner


def _styles_pdf() -> dict:
    base = getSampleStyleSheet()
    styles = {
        "marque": ParagraphStyle(
            "marque",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=15,
            textColor=OR_HEX,
            alignment=1,
            spaceBefore=6,
            spaceAfter=10,
        ),
        "objet": ParagraphStyle(
            "objet",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=17,
            textColor=BLEU_HEX,
            alignment=1,
            spaceAfter=26,
        ),
        "titre_garde": ParagraphStyle(
            "titre_garde",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=29,
            textColor=BLEU_HEX,
            alignment=1,
            spaceAfter=10,
        ),
        "client_garde": ParagraphStyle(
            "client_garde",
            parent=base["Normal"],
            fontSize=13,
            textColor=GRIS_HEX,
            alignment=1,
            spaceAfter=2,
        ),
        "mois": ParagraphStyle(
            "mois",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            textColor=GRIS_HEX,
            alignment=1,
            spaceBefore=28,
        ),
        "titre_sommaire": ParagraphStyle(
            "titre_sommaire",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=17,
            textColor=BLEU_HEX,
            spaceAfter=14,
        ),
        "section": ParagraphStyle(
            "section",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            textColor=BLEU_HEX,
            spaceBefore=14,
            spaceAfter=6,
        ),
        "corps": ParagraphStyle(
            "corps", parent=base["BodyText"], fontSize=10.5, leading=15, spaceAfter=6
        ),
        # Un élément d'énumération : retrait de bloc, puce détachée. Ce qui
        # importe est `leftIndent` — sans lui, la deuxième ligne d'un élément
        # long revenait à la marge et l'énumération se perdait.
        "puce": ParagraphStyle(
            "puce",
            parent=base["BodyText"],
            fontSize=10.5,
            leading=15,
            leftIndent=16,
            bulletIndent=4,
            spaceAfter=3,
        ),
        "sous_section": ParagraphStyle(
            "sous_section",
            parent=base["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=11,
            textColor=BLEU_HEX,
            spaceBefore=10,
            spaceAfter=4,
            leftIndent=8,
        ),
        "legende": ParagraphStyle(
            "legende",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=13,
            spaceBefore=6,
            spaceAfter=4,
        ),
        "note": ParagraphStyle(
            "note",
            parent=base["Normal"],
            fontSize=8.5,
            textColor=GRIS_HEX,
            spaceBefore=2,
            spaceAfter=8,
        ),
    }

    # Les styles de cellule. Passer par `Paragraph` plutôt que par des chaînes
    # nues est ce qui permet à un intitulé long de se replier dans sa colonne
    # au lieu de déborder sur la voisine.
    cellule = ParagraphStyle(
        "cellule", parent=base["Normal"], fontSize=8.5, leading=11
    )
    styles["cellule"] = cellule
    styles["cellule_droite"] = ParagraphStyle(
        "cellule_droite", parent=cellule, alignment=2
    )
    styles["cellule_grasse"] = ParagraphStyle(
        "cellule_grasse", parent=cellule, fontName="Helvetica-Bold"
    )
    styles["cellule_grasse_droite"] = ParagraphStyle(
        "cellule_grasse_droite", parent=styles["cellule_grasse"], alignment=2
    )
    styles["cellule_retrait"] = ParagraphStyle(
        "cellule_retrait", parent=cellule, leftIndent=10
    )
    styles["entete"] = ParagraphStyle(
        "entete", parent=cellule, fontName="Helvetica-Bold", textColor=BLEU_HEX
    )
    styles["entete_droite"] = ParagraphStyle(
        "entete_droite", parent=styles["entete"], alignment=2
    )
    return styles


def _page_de_garde_pdf(garde: Couverture, styles: dict) -> list:
    """Ce qui se pose SUR le papier à en-tête : l'objet, le titre, le mois.

    Le fond est dessiné par `_pied_de_garde`. Il porte déjà la marque, les
    filets, les coordonnées et la mention de confidentialité — les recomposer
    ici les ferait paraître deux fois.

    Le bloc descend jusqu'au panneau clair de l'en-tête, au-dessus de la bande
    d'images qui court à mi-hauteur : c'est là que le document du cabinet pose
    son titre.
    """
    elements: list = [Spacer(1, 52 * mm)]
    if garde.objet_affiche:
        elements.append(Paragraph(escape(garde.objet_affiche), styles["objet"]))
    elements.append(Paragraph(escape(garde.titre_document), styles["titre_garde"]))
    if garde.client:
        elements.append(Paragraph(escape(garde.client), styles["client_garde"]))
    if garde.reference:
        elements.append(
            Paragraph(escape(f"Réf. {garde.reference}"), styles["client_garde"])
        )
    if garde.mois:
        elements.append(Paragraph(escape(garde.mois), styles["mois"]))
    elements.append(PageBreak())
    return elements


def _sommaire_pdf(styles: dict) -> list:
    sommaire = TableOfContents()
    # Les points de conduite dès le premier niveau. Par défaut reportlab les
    # réserve aux sous-titres, et l'œil perd la ligne entre un titre court et
    # sa page, à l'autre bout de la feuille.
    sommaire.dotsMinLevel = 0
    sommaire.levelStyles = [
        ParagraphStyle(
            "somm1",
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=17,
            textColor=BLEU_HEX,
        ),
        ParagraphStyle(
            "somm2",
            fontName="Helvetica",
            fontSize=10,
            leading=15,
            leftIndent=16,
            firstLineIndent=-2,
        ),
    ]
    return [
        Paragraph(TITRE_SOMMAIRE, styles["titre_sommaire"]),
        sommaire,
        PageBreak(),
    ]


def rendre_pdf(
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None = None,
) -> bytes:
    rendues, entrees, garde = preparer(
        titre, sous_titre, sections, genere_le, couverture
    )
    styles = _styles_pdf()

    elements: list = _page_de_garde_pdf(garde, styles)
    elements += _sommaire_pdf(styles)

    def prose(section: dict, cle: str) -> list:
        return [
            Paragraph(
                f"<bullet>&bull;</bullet>{escape(texte)}" if puce else escape(texte),
                styles["puce" if puce else "corps"],
            )
            for texte, puce in paragraphes_de(section, cle)
        ]

    for section, entree in zip(rendues, entrees):
        elements.append(
            Paragraph(
                escape(entree.titre_complet),
                styles["sous_section" if entree.niveau == 2 else "section"],
            )
        )

        elements.extend(prose(section, "contenu"))

        for tableau in _tableaux_de(section):
            elements.extend(_table_pdf(tableau, styles, LARGEUR_UTILE))

        elements.extend(prose(section, "contenu_apres"))

        note = _note_de_fin(section)
        if note:
            elements.append(Paragraph(escape(note), styles["note"]))

    tampon = io.BytesIO()
    document = _DocumentRapport(
        tampon,
        pagesize=A4,
        leftMargin=MARGE,
        rightMargin=MARGE,
        topMargin=18 * mm,
        # De quoi loger le bandeau du cabinet — 11 mm de marge basse, une
        # vingtaine de millimètres de bandeau, puis le numéro de page — sans
        # que la dernière ligne du texte ne vienne mordre dessus.
        bottomMargin=38 * mm,
        title=titre,
        author="Kapi Consult",
    )
    document.multiBuild(
        elements or [Spacer(1, 1)],
        onFirstPage=_pied_de_garde(garde),
        onLaterPages=_pied_courant(garde),
    )
    return tampon.getvalue()


# --- ODT --------------------------------------------------------------------
#
# Écrit à la main plutôt qu'avec une bibliothèque. Un ODT est un zip contenant
# quelques fichiers XML ; le rapport n'a besoin que de titres, de paragraphes et
# de tables, ce qui tient en quelques dizaines de lignes. Ajouter une dépendance
# pour cela reviendrait à faire porter au déploiement le coût d'un besoin
# marginal — et l'ODT n'est demandé que par certaines administrations.

_LOGO_ODT = "Pictures/logo_kapi.png"
_BANDEAU_ODT = "Pictures/bandeau_kapi.png"

_STYLES_ODT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles
  xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
  xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
  xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
  xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
  xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
  xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
  xmlns:xlink="http://www.w3.org/1999/xlink"
  office:version="1.2">
 <!--
   La police du document. Sans cette déclaration l'ODT sortait dans la serif
   par défaut du lecteur — Liberation Serif chez LibreOffice — quand le DOCX
   et le PDF sont en Calibri et en Helvetica. Trois formats du même rapport
   n'ont pas à se ressembler de loin seulement.

   Carlito est le clone métrique de Calibri livré avec LibreOffice ; il prend
   le relais là où Calibri n'est pas installé, sans changer la mise en page.
 -->
 <office:font-face-decls>
  <style:font-face style:name="Calibri"
    svg:font-family="Calibri, Carlito, 'Liberation Sans'"
    style:font-family-generic="swiss" style:font-pitch="variable"/>
 </office:font-face-decls>
 <office:styles>
  <!-- L'interligne entre paragraphes, que le DOCX pose à 8 points. Sans lui,
       la prose de l'ODT arrivait en un seul bloc compact. -->
  <style:default-style style:family="paragraph">
   <style:paragraph-properties fo:margin-bottom="0.28cm"/>
   <style:text-properties style:font-name="Calibri" fo:font-size="11pt"/>
  </style:default-style>
  <!--
    ATTENTION À L'ORDRE. Dans un `style:style`, ODF décrit une *séquence* :
    `style:paragraph-properties` vient AVANT `style:text-properties`. Ces
    styles étaient écrits dans l'ordre inverse, et un lecteur qui applique le
    schéma jetait silencieusement tout ce qui tenait au paragraphe — le
    centrage de la page de garde, les marges, les retraits. Le texte gardait
    sa taille et sa couleur, ce qui rendait la panne difficile à voir : la
    garde s'affichait ferrée à gauche, tassée en haut de la page, sans qu'un
    seul style manque à l'appel.
  -->
  <style:style style:name="Titre" style:family="paragraph">
   <style:text-properties fo:font-size="20pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="Section" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.5cm"/>
   <style:text-properties fo:font-size="13pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="Discret" style:family="paragraph">
   <style:text-properties fo:font-size="9pt" fo:color="#6B7080"/>
  </style:style>
  <style:style style:name="SousSection" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.35cm" fo:margin-left="0.3cm"/>
   <style:text-properties fo:font-size="11pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="Legende" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.3cm" fo:margin-bottom="0.1cm"/>
   <style:text-properties fo:font-size="10pt" fo:font-weight="bold"/>
  </style:style>
  <style:style style:name="Cellule" style:family="paragraph">
   <style:text-properties fo:font-size="9pt"/>
  </style:style>
  <style:style style:name="CelluleD" style:family="paragraph">
   <style:paragraph-properties fo:text-align="end"/>
   <style:text-properties fo:font-size="9pt"/>
  </style:style>
  <style:style style:name="CelluleG" style:family="paragraph">
   <style:text-properties fo:font-size="9pt" fo:font-weight="bold"/>
  </style:style>
  <style:style style:name="CelluleGD" style:family="paragraph">
   <style:paragraph-properties fo:text-align="end"/>
   <style:text-properties fo:font-size="9pt" fo:font-weight="bold"/>
  </style:style>
  <style:style style:name="CelluleR" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="0.4cm"/>
   <style:text-properties fo:font-size="9pt"/>
  </style:style>
  <style:style style:name="Entete" style:family="paragraph">
   <style:text-properties fo:font-size="9pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="EnteteD" style:family="paragraph">
   <style:paragraph-properties fo:text-align="end"/>
   <style:text-properties fo:font-size="9pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <!-- La page de garde. -->
  <style:style style:name="GardeFilet" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.2cm" fo:margin-bottom="0cm"
     fo:border-bottom="0.06cm solid #B8892A" fo:padding-bottom="0.1cm"/>
   <style:text-properties fo:font-size="2pt"/>
  </style:style>
  <style:style style:name="GardeMarque" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-bottom="0.3cm"
     fo:border-bottom="0.06cm solid #B8892A" fo:padding-bottom="0.15cm"/>
   <style:text-properties fo:font-size="16pt" fo:font-weight="bold" fo:color="#B8892A"/>
  </style:style>
  <style:style style:name="GardeObjet" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-top="4cm"
     fo:margin-bottom="1cm"/>
   <style:text-properties fo:font-size="12pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="GardeTitre" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-bottom="0.3cm"/>
   <style:text-properties fo:font-size="24pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="GardeClient" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-bottom="0cm"/>
   <style:text-properties fo:font-size="13pt" fo:color="#6B7080"/>
  </style:style>
  <style:style style:name="GardeMois" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-top="1.2cm"/>
   <style:text-properties fo:font-size="11pt" fo:font-weight="bold" fo:color="#6B7080"/>
  </style:style>
  <style:style style:name="GardePied" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-bottom="0cm"/>
   <style:text-properties fo:font-size="7.5pt" fo:color="#6B7080"/>
  </style:style>
  <!-- La première ligne des coordonnées, qui creuse l'écart avec le titre.
       Le PDF et le DOCX posent ce bloc dans le pied de page, où il tient le
       bas de la feuille ; l'ODT n'a pas cette ressource pour une page isolée
       et s'en approche avec une marge. -->
  <style:style style:name="GardePiedDebut" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-top="5cm" fo:margin-bottom="0cm"
     fo:border-top="0.04cm solid #B8892A" fo:padding-top="0.3cm"/>
   <style:text-properties fo:font-size="7.5pt" fo:color="#6B7080"/>
  </style:style>
  <style:style style:name="Centre" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center"/>
  </style:style>
  <style:style style:name="SautDePage" style:family="paragraph">
   <style:paragraph-properties fo:break-before="page"/>
  </style:style>
  <style:style style:name="TitreSommaire" style:family="paragraph">
   <style:paragraph-properties fo:margin-bottom="0.5cm"/>
   <style:text-properties fo:font-size="17pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="Somm1" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.15cm"/>
   <style:text-properties fo:font-size="11pt" fo:font-weight="bold" fo:color="#1E2299"/>
  </style:style>
  <style:style style:name="Somm2" style:family="paragraph">
   <style:paragraph-properties fo:margin-left="0.6cm"/>
   <style:text-properties fo:font-size="10pt"/>
  </style:style>
  <style:style style:name="Pied" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center"/>
   <style:text-properties fo:font-size="8pt" fo:color="#6B7080"/>
  </style:style>
  <style:style style:name="Puce" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.05cm" fo:margin-bottom="0.05cm"/>
  </style:style>
  <style:style style:name="frBandeau" style:family="graphic">
   <style:graphic-properties style:vertical-pos="middle" style:vertical-rel="text"
     style:horizontal-pos="center" style:horizontal-rel="paragraph"/>
  </style:style>
  <!-- Une vraie liste ODF : LibreOffice tient le retrait de la seconde ligne,
       ce qu'un paragraphe commençant par un tiret ne fait pas. -->
  <text:list-style style:name="Puces">
   <text:list-level-style-bullet text:level="1" text:bullet-char="•">
    <style:list-level-properties text:space-before="0.5cm"
      text:min-label-width="0.4cm"/>
   </text:list-level-style-bullet>
  </text:list-style>
 </office:styles>
 <office:automatic-styles>
  <style:page-layout style:name="pm1">
   <style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm"
     fo:margin-top="2cm" fo:margin-bottom="2cm" fo:margin-left="2cm" fo:margin-right="2cm"/>
   <style:footer-style>
    <style:header-footer-properties fo:min-height="0.6cm" fo:margin-top="0.4cm"/>
   </style:footer-style>
  </style:page-layout>
 </office:automatic-styles>
 <office:master-styles>
  <style:master-page style:name="Standard" style:page-layout-name="pm1">
   <style:footer>
    <text:p text:style-name="Pied">__PIED__ &#8212; page
     <text:page-number text:select-page="current">1</text:page-number>
    </text:p>
    <text:p text:style-name="Pied">
     <draw:frame draw:style-name="frBandeau" text:anchor-type="as-char"
       svg:width="16cm" svg:height="1.85cm" draw:z-index="1">
      <draw:image xlink:href="__BANDEAU__" xlink:type="simple"
        xlink:show="embed" xlink:actuate="onLoad"/>
     </draw:frame>
    </text:p>
   </style:footer>
  </style:master-page>
 </office:master-styles>
</office:document-styles>
"""

# Styles de cellule : bordure partout, fond selon le genre de la ligne. Ils
# vivent dans les styles automatiques de content.xml, où ODF attend les styles
# propres à un document.
_STYLES_CELLULES_ODT = "".join(
    f'<style:style style:name="C{nom}" style:family="table-cell">'
    f'<style:table-cell-properties fo:border="0.02cm solid #C9CBD8" '
    f'fo:padding="0.08cm"'
    + (f' fo:background-color="#{fond}"' if fond else "")
    + "/></style:style>"
    for nom, fond in (
        ("Std", ""),
        ("Ent", FOND_ENTETE),
        ("Rub", FOND_RUBRIQUE),
        ("Tot", FOND_TOTAL),
    )
)

_STYLE_LOGO_ODT = (
    '<style:style style:name="frLogo" style:family="graphic">'
    '<style:graphic-properties style:vertical-pos="middle" '
    'style:vertical-rel="text" style:horizontal-pos="center" '
    'style:horizontal-rel="paragraph"/></style:style>'
)

_MANIFESTE_ODT = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest
  xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"
  manifest:version="1.2">
 <manifest:file-entry manifest:full-path="/"
   manifest:media-type="application/vnd.oasis.opendocument.text"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="__LOGO__" manifest:media-type="image/png"/>
 <manifest:file-entry manifest:full-path="__BANDEAU__" manifest:media-type="image/png"/>
</manifest:manifest>
"""


def _table_odt(tableau: dict, nom: str) -> str:
    colonnes = _colonnes(tableau)
    lignes = _lignes(tableau)
    legende = str(tableau.get("titre") or "").strip()

    morceaux = [
        f'<text:p text:style-name="Legende">{escape(l)}</text:p>'
        for l in (legende.split("\n") if legende else [])
    ]
    if not colonnes or not lignes:
        return "".join(morceaux)

    def cellule(texte: str, style_cellule: str, style_texte: str) -> str:
        return (
            f'<table:table-cell table:style-name="C{style_cellule}" '
            f'office:value-type="string">'
            f'<text:p text:style-name="{style_texte}">{escape(texte)}</text:p>'
            "</table:table-cell>"
        )

    morceaux.append(f'<table:table table:name="{nom}">')
    morceaux.append(
        f'<table:table-column table:number-columns-repeated="{len(colonnes)}"/>'
    )

    morceaux.append("<table:table-header-rows><table:table-row>")
    for colonne in colonnes:
        morceaux.append(
            cellule(
                str(colonne.get("libelle") or ""),
                "Ent",
                "EnteteD" if colonne.get("numerique") else "Entete",
            )
        )
    morceaux.append("</table:table-row></table:table-header-rows>")

    for ligne in lignes:
        genre = str(ligne.get("genre") or "")
        gras = genre in (GENRE_RUBRIQUE, GENRE_TOTAL)
        fond = {GENRE_RUBRIQUE: "Rub", GENRE_TOTAL: "Tot"}.get(genre, "Std")
        morceaux.append("<table:table-row>")
        for index, valeur in enumerate(_cellules(ligne, len(colonnes))):
            droite = bool(colonnes[index].get("numerique"))
            if genre == GENRE_DETAIL and index == 0:
                style_texte = "CelluleR"
            elif gras:
                style_texte = "CelluleGD" if droite else "CelluleG"
            else:
                style_texte = "CelluleD" if droite else "Cellule"
            morceaux.append(cellule(valeur, fond, style_texte))
        morceaux.append("</table:table-row>")

    morceaux.append("</table:table><text:p/>")
    return "".join(morceaux)


def _page_de_garde_odt(garde: Couverture) -> list[str]:
    # Le logo porte déjà le mot « Kapi Consult » : écrire NOM_CABINET dessous
    # le ferait paraître deux fois. Et ses proportions sont celles du fichier —
    # il est large, pas carré, et le forcer dans un carré l'écrasait.
    largeur = 6.0
    corps = [
        '<text:p text:style-name="Centre">'
        '<draw:frame draw:style-name="frLogo" text:anchor-type="as-char" '
        f'svg:width="{largeur:g}cm" '
        f'svg:height="{largeur * frontispice.PROPORTION_LOGO:.2f}cm" '
        'draw:z-index="0">'
        f'<draw:image xlink:href="{_LOGO_ODT}" xlink:type="simple" '
        'xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></text:p>',
        '<text:p text:style-name="GardeFilet"/>',
    ]
    if garde.objet_affiche:
        corps.append(
            f'<text:p text:style-name="GardeObjet">{escape(garde.objet_affiche)}</text:p>'
        )
    else:
        corps.append('<text:p text:style-name="GardeObjet"/>')
    corps.append(
        f'<text:p text:style-name="GardeTitre">{escape(garde.titre_document)}</text:p>'
    )
    if garde.client:
        corps.append(
            f'<text:p text:style-name="GardeClient">{escape(garde.client)}</text:p>'
        )
    if garde.reference:
        corps.append(
            f'<text:p text:style-name="GardeClient">Réf. {escape(garde.reference)}</text:p>'
        )
    if garde.mois:
        corps.append(f'<text:p text:style-name="GardeMois">{escape(garde.mois)}</text:p>')
    for rang, ligne in enumerate(frontispice.ADRESSE):
        style = "GardePiedDebut" if rang == 0 else "GardePied"
        corps.append(f'<text:p text:style-name="{style}">{escape(ligne)}</text:p>')
    corps.append(
        f'<text:p text:style-name="GardePied">{escape(frontispice.MENTION_CONFIDENTIEL)}'
        "</text:p>"
    )
    return corps


# Le gabarit d'une ligne de sommaire : le texte du titre, une tabulation à
# points de conduite, le numéro de page. LibreOffice s'en sert pour refaire la
# table ; sans lui, une actualisation produirait des lignes nues.
def _gabarit_sommaire(niveau: int) -> str:
    return (
        f'<text:table-of-content-entry-template text:outline-level="{niveau}" '
        f'text:style-name="Somm{niveau}">'
        '<text:index-entry-text/>'
        '<text:index-entry-tab-stop style:type="right" style:leader-char="."/>'
        '<text:index-entry-page-number/>'
        "</text:table-of-content-entry-template>"
    )


def _sommaire_odt(entrees: list[frontispice.Entree]) -> str:
    """Un vrai index ODF, dont le corps porte les titres dès l'ouverture.

    Ce que cet index fait, vérifié sous LibreOffice : il affiche les titres et
    leur hiérarchie, **sans numéros de page**. ODF n'a pas d'équivalent du
    `updateFields` de Word : un index ne se pagine qu'une fois actualisé par
    le lecteur (Outils › Actualiser › Index), et une conversion en PDF ne
    l'actualise pas davantage.

    On s'en tient donc à des titres justes sans pagination, plutôt qu'à des
    numéros inventés — nous ne savons pas comment LibreOffice paginera. Le
    DOCX et le PDF, eux, portent de vrais numéros ; l'ODT reste un format
    secondaire, demandé par certaines administrations.
    """
    lignes = "".join(
        f'<text:p text:style-name="Somm{e.niveau}">{escape(e.titre_complet)}</text:p>'
        for e in entrees
    )
    return (
        '<text:table-of-content text:name="Sommaire" text:protected="true">'
        '<text:table-of-content-source text:outline-level="2" '
        'text:use-outline-level="true">'
        f'<text:index-title-template text:style-name="TitreSommaire">'
        f"{escape(TITRE_SOMMAIRE)}</text:index-title-template>"
        + _gabarit_sommaire(1)
        + _gabarit_sommaire(2)
        + "</text:table-of-content-source>"
        '<text:index-body>'
        '<text:index-title text:name="Sommaire_Titre">'
        f'<text:p text:style-name="TitreSommaire">{escape(TITRE_SOMMAIRE)}</text:p>'
        "</text:index-title>"
        f"{lignes}</text:index-body></text:table-of-content>"
    )


def rendre_odt(
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None = None,
) -> bytes:
    rendues, entrees, garde = preparer(
        titre, sous_titre, sections, genere_le, couverture
    )

    corps: list[str] = _page_de_garde_odt(garde)
    corps.append('<text:p text:style-name="SautDePage"/>')
    corps.append(_sommaire_odt(entrees))
    corps.append('<text:p text:style-name="SautDePage"/>')

    for index, (section, entree) in enumerate(zip(rendues, entrees)):
        corps.append(
            f'<text:h text:outline-level="{entree.niveau}" '
            f'text:style-name="{"SousSection" if entree.niveau == 2 else "Section"}">'
            f"{escape(entree.titre_complet)}</text:h>"
        )
        def prose(section: dict, cle: str) -> list[str]:
            return [
                f'<text:list text:style-name="Puces"><text:list-item>'
                f'<text:p text:style-name="Puce">{escape(texte)}</text:p>'
                "</text:list-item></text:list>"
                if puce
                else f"<text:p>{escape(texte)}</text:p>"
                for texte, puce in paragraphes_de(section, cle)
            ]

        corps.extend(prose(section, "contenu"))
        for rang, tableau in enumerate(_tableaux_de(section)):
            corps.append(_table_odt(tableau, f"t{index}_{rang}"))
        corps.extend(prose(section, "contenu_apres"))
        note = _note_de_fin(section)
        if note:
            corps.append(f'<text:p text:style-name="Discret">{escape(note)}</text:p>')

    contenu_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<office:document-content"
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
        ' xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"'
        ' xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"'
        ' xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"'
        ' xmlns:xlink="http://www.w3.org/1999/xlink"'
        ' xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"'
        ' office:version="1.2">'
        "<office:automatic-styles>"
        + _STYLES_CELLULES_ODT
        + _STYLE_LOGO_ODT
        + "</office:automatic-styles>"
        "<office:body><office:text>" + "".join(corps) + "</office:text></office:body>"
        "</office:document-content>"
    )

    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as archive:
        # Le mimetype doit être la première entrée et rester non compressé :
        # c'est ainsi que les lecteurs reconnaissent le format sans lire le zip.
        archive.writestr(
            zipfile.ZipInfo("mimetype"),
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr("content.xml", contenu_xml)
        archive.writestr(
            "styles.xml",
            _STYLES_ODT.replace(
                "__PIED__",
                escape(f"{garde.titre_document} — {frontispice.MENTION_CONFIDENTIEL}"),
            ).replace("__BANDEAU__", _BANDEAU_ODT),
        )
        archive.writestr(_LOGO_ODT, frontispice.logo_png())
        archive.writestr(_BANDEAU_ODT, frontispice.bandeau_png())
        archive.writestr(
            "META-INF/manifest.xml",
            _MANIFESTE_ODT.replace("__LOGO__", _LOGO_ODT).replace(
                "__BANDEAU__", _BANDEAU_ODT
            ),
        )
    return tampon.getvalue()


# --- texte brut -------------------------------------------------------------


def rendre_txt(
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None = None,
) -> bytes:
    rendues, entrees, garde = preparer(
        titre, sous_titre, sections, genere_le, couverture
    )

    # La page de garde, telle qu'elle se rend sans mise en page : les mêmes
    # mentions, dans le même ordre, séparées par des filets.
    morceaux = ["=" * 72, frontispice.NOM_CABINET, "=" * 72, ""]
    if garde.objet_affiche:
        morceaux += [garde.objet_affiche, ""]
    morceaux.append(garde.titre_document)
    if garde.client:
        morceaux.append(garde.client)
    if garde.reference:
        morceaux.append(f"Réf. {garde.reference}")
    if garde.mois:
        morceaux.append(garde.mois)
    morceaux += [
        "",
        f"Établi le {genere_le.strftime('%d/%m/%Y')}",
        "",
        *frontispice.ADRESSE,
        frontispice.MENTION_CONFIDENTIEL,
        "",
        "-" * 72,
        TITRE_SOMMAIRE.upper(),
        "-" * 72,
    ]
    # Sans pagination : en texte brut il n'y a pas de page à laquelle renvoyer.
    morceaux += [
        ("    " if e.niveau == 2 else "") + e.titre_complet for e in entrees
    ]
    morceaux += ["", "=" * 72, ""]

    for section, entree in zip(rendues, entrees):
        titre_section = entree.titre_complet
        # Une sous-section se distingue par un soulignement plus discret : en
        # texte brut, c'est tout ce dont on dispose pour montrer un rang.
        if entree.niveau == 2:
            morceaux += [titre_section, "-" * len(titre_section)]
        else:
            morceaux += [titre_section.upper(), "=" * len(titre_section)]
        contenu = str(section.get("contenu") or "").strip()
        if contenu:
            morceaux += [contenu]
        # Les tableaux, alignés à l'espace : ici la chasse est fixe, c'est le
        # bon outil. Un rapport ancien n'a pas ce champ et porte son texte
        # dans `contenu`.
        aligne = str(section.get("tableaux_texte") or "").strip()
        if aligne and aligne != contenu:
            morceaux += ["", aligne] if contenu else [aligne]
        apres = str(section.get("contenu_apres") or "").strip()
        if apres:
            morceaux += ["", apres]
        note = _note_de_fin(section)
        if note:
            morceaux += ["", note]
        morceaux.append("")
    return "\n".join(morceaux).encode("utf-8")


FORMATS: dict[str, tuple[str, str]] = {
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "pdf": ("application/pdf", "pdf"),
    "odt": ("application/vnd.oasis.opendocument.text", "odt"),
    "txt": ("text/plain; charset=utf-8", "txt"),
}

_RENDUS = {
    "docx": rendre_docx,
    "pdf": rendre_pdf,
    "odt": rendre_odt,
    "txt": rendre_txt,
}


def rendre(
    format_demande: str,
    titre: str,
    sous_titre: str,
    sections: list[dict],
    genere_le: datetime,
    couverture: Couverture | None = None,
) -> tuple[bytes, str, str]:
    """Rend le rapport. Renvoie (octets, type MIME, extension)."""
    cle = format_demande.lower().lstrip(".")
    if cle not in _RENDUS:
        raise ValueError(
            f"Format inconnu : {format_demande}. Formats disponibles : "
            f"{', '.join(sorted(FORMATS))}."
        )
    mime, extension = FORMATS[cle]
    octets = _RENDUS[cle](titre, sous_titre, sections, genere_le, couverture)
    return octets, mime, extension
