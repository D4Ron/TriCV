"""End-to-end: fiche -> upload -> analysis -> ranking -> override -> export."""

from __future__ import annotations

import io
import zipfile

import pytest
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate

API = "/api/v1"


def make_cv_pdf(name: str, email: str, phone: str) -> bytes:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=name)
    doc.build([
        Paragraph(f"<b>{name}</b>", styles["Title"]),
        Paragraph(f"Email : {email}<br/>Tél : {phone}", styles["Normal"]),
        Paragraph("Adresse : 15 rue des Palmiers, Tokoin, Lomé", styles["Normal"]),
        Paragraph("Sexe : Masculin — Situation familiale : Marié", styles["Normal"]),
        Paragraph("<b>EXPERIENCE</b>", styles["Heading3"]),
        Paragraph(
            "2019-2024 Ingenieur backend senior, Orabank Togo. Conception d'une "
            "plateforme de paiement en Python et FastAPI, base PostgreSQL, "
            "deploiement Docker et integration continue.",
            styles["Normal"],
        ),
        Paragraph(
            "2016-2019 Developpeur, CIB Lome. Maintenance d'applications Django et "
            "conception de schemas relationnels.",
            styles["Normal"],
        ),
        Paragraph("<b>FORMATION</b>", styles["Heading3"]),
        Paragraph("2016 Master en Informatique, Universite de Lome", styles["Normal"]),
        Paragraph("<b>COMPETENCES</b>", styles["Heading3"]),
        Paragraph("Python, FastAPI, PostgreSQL, Docker, Git, pytest", styles["Normal"]),
    ])
    return buffer.getvalue()


async def create_session(client, auth, *, must_have_count: int = 2, criteria_count: int = 8):
    criteria = [
        {
            "name": f"Critère {i + 1}",
            "description": f"Ce que nous attendons pour le critère {i + 1}.",
            "weight": (i % 10) + 1,
            "is_must_have": i < must_have_count,
            "display_order": i,
        }
        for i in range(criteria_count)
    ]
    response = await client.post(
        f"{API}/sessions",
        headers=auth,
        json={
            "title": "Développeur Backend Senior",
            "position": "Développeur Backend Senior",
            "description": "Renfort sur une plateforme de paiement.",
            "department": "Technologie",
            "language": "fr",
            "status": "OPEN",
            "score_threshold": 60,
            "criteria": criteria,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


async def upload_cv(client, auth, session_id: str, name: str, email: str, phone: str):
    return await client.post(
        f"{API}/sessions/{session_id}/candidates",
        headers=auth,
        files={"files": (f"{name}.pdf", make_cv_pdf(name, email, phone), "application/pdf")},
        data={"full_name": name, "email": email},
    )


# --- foundation -------------------------------------------------------------


async def test_health(client):
    response = await client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


async def test_auth_is_required(client):
    assert (await client.get(f"{API}/sessions")).status_code == 401


async def test_login_and_me(client, auth):
    response = await client.get(f"{API}/auth/me", headers=auth)
    assert response.status_code == 200
    assert response.json()["role"] == "ADMIN"


async def test_bad_password_is_rejected(client, admin_token):
    response = await client.post(
        f"{API}/auth/login", json={"email": "admin@tricv.example", "password": "wrong"}
    )
    assert response.status_code == 401


# --- sessions ---------------------------------------------------------------


async def test_create_and_fetch_session(client, auth):
    session = await create_session(client, auth)
    assert len(session["criteria"]) == 8
    assert len(session["public_key"]) == 32
    assert session["criteria_locked"] is False

    detail = await client.get(f"{API}/sessions/{session['id']}", headers=auth)
    assert detail.status_code == 200
    assert detail.json()["counts"]["total"] == 0


async def test_a_session_cannot_open_without_criteria(client, auth):
    response = await client.post(
        f"{API}/sessions",
        headers=auth,
        json={"title": "X", "position": "X", "status": "OPEN", "criteria": []},
    )
    assert response.status_code == 422


async def test_duplicate_session_copies_criteria_but_no_candidates(client, auth):
    original = await create_session(client, auth)
    response = await client.post(
        f"{API}/sessions/{original['id']}/duplicate", headers=auth, json={}
    )
    assert response.status_code == 201
    copy = response.json()
    assert copy["status"] == "DRAFT"
    assert copy["public_key"] != original["public_key"]
    assert [c["name"] for c in copy["criteria"]] == [c["name"] for c in original["criteria"]]
    assert copy["counts"]["total"] == 0


# --- the analysis pipeline --------------------------------------------------


async def test_upload_analyses_and_ranks(client, auth, stub_provider):
    session = await create_session(client, auth)

    response = await upload_cv(
        client, auth, session["id"], "Amevi KOSSI", "amevi.kossi@example.tg", "+228 90 12 34 56"
    )
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["accepted"] == 1

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    assert listing.status_code == 200
    items = listing.json()["items"]
    assert len(items) == 1

    candidate = items[0]
    assert candidate["analysis_status"] == "ANALYZED", candidate.get("analysis_error")
    assert candidate["rank"] == 1
    assert candidate["ai_score"] == 75.0
    assert candidate["ai_recommendation"] == "FIT"
    assert candidate["redaction_applied"] is True


async def test_the_provider_payload_contains_no_identifiers(client, auth, stub_provider):
    """Acceptance: with PII_REDACTION=true the exact payload sent can be
    inspected and contains no name, email, phone or address."""
    session = await create_session(client, auth)
    await upload_cv(
        client, auth, session["id"], "Amevi KOSSI", "amevi.kossi@example.tg", "+228 90 12 34 56"
    )

    assert len(stub_provider.sent_payloads) == 1
    payload = stub_provider.sent_payloads[0]

    assert payload.redacted is True
    assert payload.file_bytes is None, "the original file must not leave the system"
    lowered = payload.text.lower()
    for identifier in ("kossi", "amevi", "amevi.kossi@example.tg", "90 12 34 56", "palmiers"):
        assert identifier not in lowered

    # ...and demographics are gone from the payload but visible to HR.
    assert "masculin" not in lowered
    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    detail = await client.get(
        f"{API}/candidates/{listing.json()['items'][0]['id']}", headers=auth
    )
    assert detail.json()["demographics"]["gender"] == "Masculin"


async def test_the_model_cannot_leak_identity_back_into_the_record(client, auth, stub_provider):
    """The stub returns a name and email. On the redacted path the provider
    never saw them, so anything it returns is invention and must be ignored."""
    session = await create_session(client, auth)
    await upload_cv(
        client, auth, session["id"], "Amevi KOSSI", "amevi.kossi@example.tg", "+228 90 12 34 56"
    )

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate = listing.json()["items"][0]
    assert candidate["full_name"] == "Amevi KOSSI"
    assert candidate["email"] == "amevi.kossi@example.tg"


async def test_missing_must_have_is_marked_not_fit_with_the_item_named(
    client, auth, stub_provider
):
    session = await create_session(client, auth, must_have_count=2)
    must_have_name = session["criteria"][1]["name"]
    must_have_id = session["criteria"][1]["id"]

    stub_provider.score_for = lambda c: 20.0 if c.id == must_have_id else 95.0

    await upload_cv(client, auth, session["id"], "Faible Profil", "faible@example.tg", "+22890000000")

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate = listing.json()["items"][0]

    assert candidate["ai_recommendation"] == "NOT_FIT"
    assert candidate["missing_must_haves"] == [must_have_name]
    # The score is still stored, so HR sees how close they were.
    assert candidate["ai_score"] > 60


async def test_criteria_lock_once_a_candidate_is_analysed(client, auth, stub_provider):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "A B", "ab@example.tg", "+22890000001")

    response = await client.patch(
        f"{API}/sessions/{session['id']}",
        headers=auth,
        json={"criteria": [{"name": "Nouveau", "weight": 5, "is_must_have": False}]},
    )
    assert response.status_code == 409
    assert "duplicate" in response.json()["detail"].lower()

    # Non-criteria fields still update fine.
    assert (
        await client.patch(
            f"{API}/sessions/{session['id']}", headers=auth, json={"score_threshold": 70}
        )
    ).status_code == 200


async def test_unreadable_cv_fails_with_a_clear_message(client, auth, stub_provider):
    """A scan has no extractable text — HR is told the file is not machine-readable."""
    session = await create_session(client, auth)

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4).build(
        [Paragraph(".", getSampleStyleSheet()["Normal"])]
    )

    response = await client.post(
        f"{API}/sessions/{session['id']}/candidates",
        headers=auth,
        files={"files": ("scan.pdf", buffer.getvalue(), "application/pdf")},
    )
    assert response.status_code == 202

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate = listing.json()["items"][0]
    assert candidate["analysis_status"] == "FAILED"
    assert "machine-readable" in candidate["analysis_error"]


async def test_non_document_uploads_are_rejected_on_magic_bytes(client, auth):
    session = await create_session(client, auth)
    response = await client.post(
        f"{API}/sessions/{session['id']}/candidates",
        headers=auth,
        # A .pdf extension over PNG bytes: the extension is not trusted.
        files={"files": ("evil.pdf", b"\x89PNG\r\n\x1a\n" + b"0" * 400, "application/pdf")},
    )
    assert response.status_code == 202
    assert response.json()["accepted"] == 0
    # Le refus nomme le fichier : sur un lot, « un fichier a été refusé » sans
    # dire lequel n'aide personne.
    erreur = response.json()["results"][0]["error"]
    assert "evil.pdf" in erreur and "Word" in erreur


async def test_duplicate_files_are_flagged_not_blocked(client, auth, stub_provider):
    session = await create_session(client, auth)
    pdf = make_cv_pdf("Amevi KOSSI", "amevi@example.tg", "+22890123456")

    for _ in range(2):
        response = await client.post(
            f"{API}/sessions/{session['id']}/candidates",
            headers=auth,
            files={"files": ("cv.pdf", pdf, "application/pdf")},
        )
        assert response.status_code == 202

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    items = listing.json()["items"]
    assert len(items) == 2, "both submissions are kept — HR decides"
    assert all(item["duplicates"] for item in items)
    assert items[0]["duplicates"][0]["reason"] == "identical_file"


# --- HR overrides -----------------------------------------------------------


async def test_manual_score_overrides_the_ranking(client, auth, stub_provider):
    session = await create_session(client, auth)
    stub_provider.score_for = lambda c: 90.0
    await upload_cv(client, auth, session["id"], "Fort Profil", "fort@example.tg", "+22890000002")
    stub_provider.score_for = lambda c: 50.0
    await upload_cv(client, auth, session["id"], "Moyen Profil", "moyen@example.tg", "+22890000003")

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    items = listing.json()["items"]
    assert items[0]["full_name"] == "Fort Profil"

    weaker = next(i for i in items if i["full_name"] == "Moyen Profil")
    patched = await client.patch(
        f"{API}/candidates/{weaker['id']}",
        headers=auth,
        json={"manual_score": 99, "hr_status": "SHORTLISTED", "hr_notes": "Excellent entretien."},
    )
    assert patched.status_code == 200
    assert patched.json()["effective_score"] == 99.0

    reordered = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    assert reordered.json()["items"][0]["full_name"] == "Moyen Profil"


async def test_overrides_are_audited(client, auth, stub_provider):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "A B", "ab@example.tg", "+22890000004")
    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate_id = listing.json()["items"][0]["id"]

    await client.patch(
        f"{API}/candidates/{candidate_id}", headers=auth, json={"hr_status": "REJECTED"}
    )

    audit = await client.get(f"{API}/sessions/{session['id']}/audit", headers=auth)
    actions = [entry["action"] for entry in audit.json()]
    assert "candidate.override" in actions


async def test_full_document_reanalysis_is_audited_and_bypasses_redaction(
    client, auth, stub_provider
):
    session = await create_session(client, auth)
    await upload_cv(
        client, auth, session["id"], "Amevi KOSSI", "amevi@example.tg", "+22890123456"
    )
    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate_id = listing.json()["items"][0]["id"]

    response = await client.post(
        f"{API}/candidates/{candidate_id}/reanalyze",
        headers=auth,
        json={"full_document": True},
    )
    assert response.status_code == 202

    # The second payload carries the original file, not redacted text.
    assert len(stub_provider.sent_payloads) == 2
    second = stub_provider.sent_payloads[1]
    assert second.file_bytes is not None
    assert second.redacted is False

    audit = await client.get(f"{API}/sessions/{session['id']}/audit", headers=auth)
    entry = next(e for e in audit.json() if e["action"] == "candidate.reanalyze_full_document")
    assert entry["details"]["redaction_bypassed"] is True
    assert entry["user_name"] == "Test Admin"


async def test_reanalysis_replaces_scores_instead_of_colliding(client, auth, stub_provider):
    """Re-analysis rewrites the criterion scores. They are unique per
    (candidate, criterion), so the old rows must be gone before the new ones
    are written — otherwise every re-run dies on the unique constraint."""
    session = await create_session(client, auth, criteria_count=4)
    await upload_cv(client, auth, session["id"], "A B", "ab@example.tg", "+22890000011")

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate_id = listing.json()["items"][0]["id"]

    first = await client.get(f"{API}/candidates/{candidate_id}", headers=auth)
    assert first.json()["analysis_status"] == "ANALYZED"
    assert len(first.json()["criterion_scores"]) == 4

    # Second pass, with different scores so we can prove they were replaced.
    stub_provider.score_for = lambda c: 42.0
    response = await client.post(
        f"{API}/candidates/{candidate_id}/reanalyze", headers=auth, json={"full_document": False}
    )
    assert response.status_code == 202

    after = await client.get(f"{API}/candidates/{candidate_id}", headers=auth)
    body = after.json()
    assert body["analysis_status"] == "ANALYZED", body["analysis_error"]
    assert len(body["criterion_scores"]) == 4, "scores were duplicated or dropped"
    assert all(s["score"] == 42.0 for s in body["criterion_scores"])
    assert body["ai_score"] == 42.0


# --- public intake ----------------------------------------------------------


async def test_public_session_and_apply(client, auth, stub_provider):
    session = await create_session(client, auth)
    key = session["public_key"]

    public = await client.get(f"{API}/public/session/{key}")
    assert public.status_code == 200
    assert public.json()["accepts_applications"] is True
    # A candidate must never see scores or other applicants.
    assert "score_threshold" not in public.json()
    assert "criteria" not in public.json()

    applied = await client.post(
        f"{API}/public/apply/{key}",
        files={"cv": ("cv.pdf", make_cv_pdf("Afiwa MENSAH", "afiwa@example.tg", "+22891112233"),
                      "application/pdf")},
        data={"full_name": "Afiwa MENSAH", "email": "afiwa@example.tg", "phone": "+228 91 11 22 33"},
    )
    assert applied.status_code == 202, applied.text
    assert "score" not in applied.text.lower()

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    candidate = listing.json()["items"][0]
    assert candidate["source"] == "PUBLIC_FORM"
    assert candidate["full_name"] == "Afiwa MENSAH"


async def test_closed_session_stops_accepting_applications(client, auth):
    session = await create_session(client, auth)
    await client.post(f"{API}/sessions/{session['id']}/close", headers=auth)

    response = await client.post(
        f"{API}/public/apply/{session['public_key']}",
        files={"cv": ("cv.pdf", make_cv_pdf("X Y", "x@example.tg", "+22890000005"),
                      "application/pdf")},
        data={"full_name": "X Y", "email": "x@example.tg"},
    )
    assert response.status_code == 409


async def test_unknown_public_key_is_a_404(client):
    assert (await client.get(f"{API}/public/session/nope")).status_code == 404


async def test_careers_index_lists_only_open_public_roles(client, auth):
    """The public careers index is unauthenticated, so it must show open roles
    and nothing else — no drafts, no closed roles, no opted-out ones."""
    open_session = await create_session(client, auth)

    draft = await client.post(
        f"{API}/sessions",
        headers=auth,
        json={
            "title": "Brouillon interne",
            "position": "Poste en préparation",
            "status": "DRAFT",
            "criteria": [{"name": "C", "weight": 5, "is_must_have": False}],
        },
    )
    assert draft.status_code == 201

    closed = await create_session(client, auth)
    await client.post(f"{API}/sessions/{closed['id']}/close", headers=auth)

    opted_out = await create_session(client, auth)
    await client.patch(
        f"{API}/sessions/{opted_out['id']}",
        headers=auth,
        json={"accepts_public_applications": False},
    )

    response = await client.get(f"{API}/public/roles")
    assert response.status_code == 200
    roles = response.json()

    keys = {role["public_key"] for role in roles}
    assert open_session["public_key"] in keys
    assert closed["public_key"] not in keys
    assert opted_out["public_key"] not in keys
    assert "Poste en préparation" not in [role["position"] for role in roles]


async def test_careers_index_never_exposes_scoring_internals(client, auth, stub_provider):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "A B", "ab@example.tg", "+22890000010")

    roles = (await client.get(f"{API}/public/roles")).json()
    role = next(r for r in roles if r["public_key"] == session["public_key"])

    assert set(role) == {
        "public_key", "title", "position", "department", "description", "language", "posted_at",
    }
    for leaked in ("criteria", "counts", "score_threshold", "candidates", "id"):
        assert leaked not in role


# --- stats and structure-fiche ---------------------------------------------


async def test_stats(client, auth, stub_provider):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "A B", "ab@example.tg", "+22890000006")

    stats = await client.get(f"{API}/sessions/{session['id']}/stats", headers=auth)
    body = stats.json()
    assert body["total_candidates"] == 1
    assert body["by_recommendation"]["FIT"] == 1
    assert body["average_score"] == 75.0
    assert len(body["distribution"]) == 5


async def test_structure_fiche_returns_drafts_without_persisting(client, auth, stub_provider):
    response = await client.post(
        f"{API}/sessions/structure-fiche",
        headers=auth,
        json={
            "raw_text": "Nous recherchons un développeur backend Python avec au moins "
                        "trois ans d'expérience et une bonne maîtrise de PostgreSQL.",
            "language": "fr",
        },
    )
    assert response.status_code == 200
    criteria = response.json()["criteria"]
    assert len(criteria) == 2
    assert criteria[0]["is_must_have"] is True

    # Nothing was written.
    assert (await client.get(f"{API}/sessions", headers=auth)).json() == []


async def test_settings_endpoint_surfaces_the_privacy_posture(client, auth):
    response = await client.get(f"{API}/settings", headers=auth)
    body = response.json()
    assert body["pii_redaction"] is True
    assert body["redact_demographics"] is True


# --- exports ----------------------------------------------------------------


@pytest.mark.parametrize(
    "fmt,magic",
    [("pdf", b"%PDF-"), ("xlsx", b"PK\x03\x04"), ("docx", b"PK\x03\x04")],
)
async def test_exports_produce_valid_files(client, auth, stub_provider, fmt, magic):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "Amevi KOSSI", "amevi@example.tg", "+22890123456")

    response = await client.get(
        f"{API}/sessions/{session['id']}/export?format={fmt}", headers=auth
    )
    assert response.status_code == 200, response.text
    assert response.content.startswith(magic)
    assert len(response.content) > 2000
    assert "attachment" in response.headers["content-disposition"]


async def test_excel_export_opens_and_has_the_three_sheets(client, auth, stub_provider):
    from openpyxl import load_workbook

    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "Amevi KOSSI", "amevi@example.tg", "+22890123456")

    response = await client.get(
        f"{API}/sessions/{session['id']}/export?format=xlsx", headers=auth
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Classement", "Critères", "Synthèse"]

    ranking = workbook["Classement"]
    assert ranking.freeze_panes == "C2"
    assert ranking.auto_filter.ref is not None
    assert ranking.cell(row=2, column=2).value == "Amevi KOSSI"


async def test_docx_export_opens(client, auth, stub_provider):
    from docx import Document

    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "Amevi KOSSI", "amevi@example.tg", "+22890123456")

    response = await client.get(
        f"{API}/sessions/{session['id']}/export?format=docx", headers=auth
    )
    document = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Rapport de présélection" in text


async def test_zip_bundle_contains_correctly_named_cvs(client, auth, stub_provider):
    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "Amevi KOSSI", "amevi@example.tg", "+22890123456")
    await upload_cv(client, auth, session["id"], "Afiwa MENSAH", "afiwa@example.tg", "+22891112233")

    response = await client.get(
        f"{API}/sessions/{session['id']}/export?format=pdf&include_cvs=true", headers=auth
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"

    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.testzip() is None

    names = archive.namelist()
    cvs = sorted(n for n in names if n.startswith("CVs/"))
    assert len(cvs) == 2
    assert cvs[0] == "CVs/01_KOSSI_Amevi.pdf" or cvs[0] == "CVs/01_MENSAH_Afiwa.pdf"
    assert all(archive.read(n).startswith(b"%PDF-") for n in cvs)
    assert any(n.endswith(".pdf") and not n.startswith("CVs/") for n in names)


async def test_export_scope_shortlisted(client, auth, stub_provider):
    from openpyxl import load_workbook

    session = await create_session(client, auth)
    await upload_cv(client, auth, session["id"], "Retenu Profil", "a@example.tg", "+22890000007")
    await upload_cv(client, auth, session["id"], "Ecarte Profil", "b@example.tg", "+22890000008")

    listing = await client.get(f"{API}/sessions/{session['id']}/candidates", headers=auth)
    keep = next(i for i in listing.json()["items"] if i["full_name"] == "Retenu Profil")
    await client.patch(
        f"{API}/candidates/{keep['id']}", headers=auth, json={"hr_status": "SHORTLISTED"}
    )

    response = await client.get(
        f"{API}/sessions/{session['id']}/export?format=xlsx&scope=shortlisted", headers=auth
    )
    sheet = load_workbook(io.BytesIO(response.content))["Classement"]
    names = [sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)]
    assert names == ["Retenu Profil"]
