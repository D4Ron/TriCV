"""Purge des fichiers d'un mandat archivé.

Le point à tenir : purger efface les documents et rien d'autre. Ce qui justifie
une décision des mois plus tard — le parcours saisi, la note et son détail, les
motifs d'élimination — ne dépend pas du PDF, et une grille recalculée après
purge doit donner exactement le même résultat.
"""

from __future__ import annotations

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import PieceCandidature

pytestmark = pytest.mark.anyio

API = "/api/v1"
PDF = b"%PDF-1.4\n" + b"0" * 4000


async def monter(client, auth) -> dict:
    """Un mandat avec un poste, un avis et un dossier déposé."""
    ids = {}
    ids["client"] = (
        await client.post(f"{API}/clients", json={"nom": "Client"}, headers=auth)
    ).json()["id"]
    ids["mandat"] = (
        await client.post(
            f"{API}/mandats",
            json={"client_id": ids["client"], "intitule": "Mandat"},
            headers=auth,
        )
    ).json()["id"]
    ids["poste"] = (
        await client.post(
            f"{API}/mandats/{ids['mandat']}/postes",
            json={"intitule": "Poste", "niveau_min": 3, "pieces_requises": ["CV"]},
            headers=auth,
        )
    ).json()["id"]
    ids["avis"] = (
        await client.post(
            f"{API}/postes/{ids['poste']}/avis",
            json={"date_cloture": "2026-12-31"},
            headers=auth,
        )
    ).json()["id"]

    depot = await client.post(
        f"{API}/postes/{ids['poste']}/candidatures/depot-multiple",
        files=[("fichiers", ("cv.pdf", PDF, "application/pdf"))],
        headers=auth,
    )
    assert depot.status_code == 201, depot.text
    return ids


async def test_purger_un_mandat_en_cours_est_refuse(client, auth):
    """Tant que le recrutement tourne, les pièces servent encore."""
    ids = await monter(client, auth)

    refus = await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)
    assert refus.status_code == 409
    assert "archivez" in refus.json()["detail"].lower()

    # Rien n'a bougé.
    async with SessionLocal() as db:
        pieces = (await db.execute(select(PieceCandidature))).scalars().all()
        assert all(p.chemin_stockage is not None for p in pieces)


async def test_l_estimation_ne_supprime_rien(client, auth):
    ids = await monter(client, auth)

    estimation = (await client.get(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)).json()
    assert estimation["fichiers"] == 1
    assert estimation["candidatures"] == 1

    async with SessionLocal() as db:
        pieces = (await db.execute(select(PieceCandidature))).scalars().all()
        assert all(p.chemin_stockage is not None for p in pieces)


async def test_la_purge_efface_les_fichiers_et_garde_les_dossiers(client, auth):
    ids = await monter(client, auth)
    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)

    resultat = await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)
    assert resultat.status_code == 200, resultat.text
    assert resultat.json()["fichiers"] == 1

    async with SessionLocal() as db:
        pieces = (await db.execute(select(PieceCandidature))).scalars().all()
        # La ligne survit : on sait toujours quelle pièce avait été reçue.
        assert len(pieces) == 1
        assert pieces[0].chemin_stockage is None
        assert pieces[0].purge_le is not None
        assert pieces[0].taille_octets == len(PDF)

    # Le dossier reste consultable.
    candidatures = await client.get(
        f"{API}/postes/{ids['poste']}/candidatures", headers=auth
    )
    assert candidatures.status_code == 200
    assert candidatures.json()["total"] == 1


async def test_la_grille_est_identique_avant_et_apres_purge(client, auth):
    """La complétude se juge sur la pièce reçue, pas sur le fichier conservé."""
    ids = await monter(client, auth)

    await client.post(f"{API}/postes/{ids['poste']}/evaluer", headers=auth)
    avant = (await client.get(f"{API}/postes/{ids['poste']}/grille", headers=auth)).json()

    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)
    await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)

    await client.post(f"{API}/postes/{ids['poste']}/evaluer", headers=auth)
    apres = (await client.get(f"{API}/postes/{ids['poste']}/grille", headers=auth)).json()

    for champ in (
        "nombre_candidatures",
        "nombre_preselectionnes",
        "nombre_elimines",
        "nombre_a_verifier",
        "total_max",
        "seuil",
    ):
        assert avant[champ] == apres[champ], champ


async def test_purger_deux_fois_ne_compte_pas_deux_fois(client, auth):
    ids = await monter(client, auth)
    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)

    premier = (await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)).json()
    second = (await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)).json()

    assert premier["fichiers"] == 1
    assert second["fichiers"] == 0


async def test_la_purge_est_journalisee(client, auth):
    from app.models import AuditLog

    ids = await monter(client, auth)
    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)
    await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)

    async with SessionLocal() as db:
        actions = [a.action for a in (await db.execute(select(AuditLog))).scalars()]
        assert "mandat.purge" in actions



# --- la suppression, et les fichiers qu'elle laissait derrière elle ----------
#
# La cascade SQL emporte les lignes — pièces, postes, candidatures — mais pas
# ce qu'elles désignent sur le disque. Un mandat supprimé laissait donc tous
# ses CV et sa fiche de poste dans le stockage, sans plus aucune ligne pour
# dire à quoi ils correspondaient : ni retrouvables, ni effaçables, et des
# données personnelles qu'on croyait supprimées avec le mandat.
#
# Les tests comparent les fichiers **de ce mandat-ci** : le dépôt est commun à
# toute la session d'essai, et le voir vide ne prouverait rien.


def _fichiers_du_depot() -> set[str]:
    from pathlib import Path

    from app.config import settings

    racine = Path(settings.storage_path)
    return {str(p) for p in racine.rglob("*") if p.is_file()} if racine.exists() else set()


def _fiche_docx() -> bytes:
    """Une fiche de poste minimale, mais un vrai DOCX : le lecteur l'ouvre."""
    import io

    import docx

    document = docx.Document()
    document.add_paragraph("Intitulé du poste : Comptable")
    document.add_paragraph("Profil requis")
    document.add_paragraph("Diplôme de niveau BAC+3 au minimum.")
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


async def test_supprimer_un_mandat_efface_aussi_ses_fichiers(client, auth):
    avant = _fichiers_du_depot()
    ids = await monter(client, auth)
    siens = _fichiers_du_depot() - avant
    assert siens, "le dépôt doit contenir le CV déposé"

    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)
    reponse = await client.delete(
        f"{API}/mandats/{ids['mandat']}?confirmer=true", headers=auth
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["fichiers_supprimes"] >= 1

    restants = siens & _fichiers_du_depot()
    assert not restants, f"fichiers orphelins : {restants}"


async def test_supprimer_un_client_efface_les_fichiers_de_ses_mandats(client, auth):
    avant = _fichiers_du_depot()
    ids = await monter(client, auth)
    siens = _fichiers_du_depot() - avant
    assert siens

    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)
    await client.post(f"{API}/clients/{ids['client']}/archiver", headers=auth)
    reponse = await client.delete(
        f"{API}/clients/{ids['client']}?confirmer=true", headers=auth
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["fichiers_supprimes"] >= 1

    assert not (siens & _fichiers_du_depot())


async def test_la_fiche_de_poste_part_a_la_purge(client, auth):
    """Le document du client est un fichier comme un autre.

    Son texte relevé, lui, reste : c'est ce que relit la rédaction d'un avis,
    et il ne pèse rien.
    """
    ids = await monter(client, auth)
    depot = await client.post(
        f"{API}/postes/{ids['poste']}/fiche",
        files=[
            (
                "fichier",
                (
                    "fiche.docx",
                    _fiche_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )
        ],
        data={"proposer": "false"},
        headers=auth,
    )
    assert depot.status_code == 200, depot.text

    await client.post(f"{API}/mandats/{ids['mandat']}/archiver", headers=auth)
    resultat = (
        await client.post(f"{API}/mandats/{ids['mandat']}/purge", headers=auth)
    ).json()
    assert resultat["fiches"] >= 1

    from app.models import Poste

    async with SessionLocal() as db:
        poste = await db.get(Poste, ids["poste"])
        assert poste.fiche_chemin is None
        # Ce qu'on a tiré du document survit à son fichier.
        assert poste.fiche_nom_fichier == "fiche.docx"
        assert poste.fiche_texte and "Comptable" in poste.fiche_texte
