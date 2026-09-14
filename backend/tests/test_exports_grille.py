"""Les quatre documents tableur du poste, exportables séparément.

Un classeur unique de quatre onglets obligeait à envoyer plus que ce qu'on
voulait montrer : le classement part au client, le tableau d'élimination se
produit à un candidat qui conteste, et les deux ne devraient pas voyager
ensemble.

La règle que ces tests protègent : **une feuille exportée seule est mot pour
mot celle du classeur complet**. Deux chemins de production finiraient par
diverger, et l'écart ne se découvrirait qu'une fois le fichier chez le client.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.anyio

API = "/api/v1"


async def monter_poste_note(client, auth) -> str:
    client_id = (
        await client.post(f"{API}/clients", json={"nom": "Dogta-Lafiè"}, headers=auth)
    ).json()["id"]
    mandat_id = (
        await client.post(
            f"{API}/mandats", json={"client_id": client_id, "intitule": "Direction"}, headers=auth
        )
    ).json()["id"]
    poste_id = (
        await client.post(
            f"{API}/mandats/{mandat_id}/postes",
            json={
                "intitule": "Chef comptable",
                "niveau_min": 4,
                "domaines_acceptes": ["comptabilite"],
                "annees_experience_min": 5,
                "annees_experience_specifique_min": 3,
                "domaines_experience": ["comptabilite"],
                "pieces_requises": ["CV"],
            },
            headers=auth,
        )
    ).json()["id"]

    # Un dossier qui sera écarté : le tableau d'élimination doit avoir un objet.
    await client.post(
        f"{API}/postes/{poste_id}/candidatures",
        json={
            "candidat": {"nom": "Koumako", "prenom": "Edem", "email": "e.k@example.tg"},
            "pieces": [],
        },
        headers=auth,
    )
    return poste_id


async def classeur(client, auth, poste_id: str, type_grille: str | None = None):
    url = f"{API}/postes/{poste_id}/grille.xlsx"
    if type_grille:
        url += f"?type_grille={type_grille}"
    reponse = await client.get(url, headers=auth)
    assert reponse.status_code == 200, reponse.text
    return reponse, load_workbook(io.BytesIO(reponse.content))


async def test_le_defaut_reste_le_classeur_complet(client, auth):
    """Un lien existant continue de rendre ce qu'il rendait."""
    poste_id = await monter_poste_note(client, auth)
    reponse, wb = await classeur(client, auth, poste_id)

    assert "Grille de présélection" in wb.sheetnames
    assert "Tableau d'élimination" in wb.sheetnames
    assert "Synthèse" in wb.sheetnames
    assert "dossier-complet" in reponse.headers["content-disposition"]


@pytest.mark.parametrize(
    ("type_grille", "feuille", "fichier"),
    [
        ("PRESELECTION", "Grille de présélection", "grille-preselection"),
        ("ELIMINATION", "Tableau d'élimination", "tableau-elimination"),
        ("ENTRETIENS", "Entretiens", "entretiens"),
        ("SYNTHESE", "Synthèse", "synthese"),
    ],
)
async def test_chaque_document_s_exporte_seul(
    client, auth, type_grille: str, feuille: str, fichier: str
):
    poste_id = await monter_poste_note(client, auth)
    reponse, wb = await classeur(client, auth, poste_id, type_grille)

    assert wb.sheetnames == [feuille], "un document seul, sans onglet parasite"
    # Jamais l'onglet « Sheet » qu'openpyxl crée d'office : le lecteur
    # ouvrirait un classeur sur une page vide.
    assert "Sheet" not in wb.sheetnames
    assert fichier in reponse.headers["content-disposition"]


async def test_une_feuille_seule_est_celle_du_classeur_complet(client, auth):
    """Deux chemins de production divergeraient ; il n'y en a qu'un."""
    poste_id = await monter_poste_note(client, auth)
    _, complet = await classeur(client, auth, poste_id)
    _, seul = await classeur(client, auth, poste_id, "ELIMINATION")

    a = [
        [c.value for c in ligne]
        for ligne in complet["Tableau d'élimination"].iter_rows(max_row=12)
    ]
    b = [
        [c.value for c in ligne]
        for ligne in seul["Tableau d'élimination"].iter_rows(max_row=12)
    ]
    assert a == b


async def test_les_entretiens_demandes_seuls_le_disent_quand_il_n_y_en_a_pas(client, auth):
    """Un fichier sans feuille ne serait pas une réponse ; une feuille qui dit
    « aucun entretien saisi » en est une."""
    poste_id = await monter_poste_note(client, auth)
    _, wb = await classeur(client, auth, poste_id, "ENTRETIENS")

    textes = [
        str(c.value)
        for ligne in wb["Entretiens"].iter_rows(max_row=8)
        for c in ligne
        if c.value
    ]
    assert any("Aucun entretien" in t for t in textes)


async def test_un_type_inconnu_est_refuse(client, auth):
    poste_id = await monter_poste_note(client, auth)
    reponse = await client.get(
        f"{API}/postes/{poste_id}/grille.xlsx?type_grille=BULLETIN", headers=auth
    )
    assert reponse.status_code == 422
