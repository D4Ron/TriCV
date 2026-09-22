"""Le rapport remis s'ouvre sur une page de garde et un sommaire.

Le document que Kapi Consult remet — relevé sur « Rapport des entretiens -
WAPP 2025 » — ne commence pas à « Introduction ». Il commence par une page à
l'en-tête du cabinet, puis par un sommaire, et ses titres sont numérotés. Les
exports produisaient le corps seul : un fichier qui s'ouvre sur « Introduction »
n'est pas celui que le client attend.

Ce fichier vérifie donc trois choses. Que la numérotation suit la règle du
document. Que chacun des quatre formats porte la page de garde. Et que chacun
dresse son sommaire avec ses propres moyens — un champ TOC pour Word, un index
ODF pour LibreOffice, une table calculée en deux passes pour le PDF, une liste
pour le texte brut.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime

import pytest
from docx import Document

from app.services.exports import frontispice
from app.services.exports import rapport as export

QUAND = datetime(2025, 5, 12)

GARDE = frontispice.Couverture(
    titre="Rapport final de recrutement",
    objet="Recrutement de deux (2) postes pour le WAPP",
    client="WAPP — West African Power Pool",
    reference="WAPP/2025/DRH-07",
    mois="Mai 2025",
)


def section(code: str, titre: str, niveau: int = 1, contenu: str = "du texte") -> dict:
    return {
        "code": code,
        "titre": titre,
        "contenu": contenu,
        "niveau": niveau,
        "origine": "REDIGEE",
    }


def trame() -> list[dict]:
    """La trame du cabinet en réduction : « Méthodologie » n'a pas de texte à
    elle, elle porte ses deux sous-sections."""
    porteuse = section("METHODOLOGIE", "Méthodologie", contenu="")
    porteuse["porteur"] = True
    return [
        section("INTRODUCTION", "Introduction"),
        section("DEMARCHE", "Démarche"),
        porteuse,
        section("METHODE_PRESELECTION", "Présélection", niveau=2),
        section("CRITERES_ELIMINATOIRES", "Critères éliminatoires", niveau=2),
        section("RESULTATS_PRESELECTION", "Résultats de la présélection"),
    ]


def rendre(fonction, sections=None):
    return fonction(
        "Rapport final de recrutement",
        "WAPP",
        trame() if sections is None else sections,
        QUAND,
        GARDE,
    )


# --- la numérotation --------------------------------------------------------


def test_les_sections_sont_numerotees_comme_dans_le_document_remis():
    """Romain pour les sections, décimal sous elles, rien pour l'introduction."""
    rendues, entrees, _ = export.preparer("T", "C", trame(), QUAND, GARDE)
    assert [e.titre_complet for e in entrees] == [
        "Introduction",
        "I. Démarche",
        "II. Méthodologie",
        "2.1. Présélection",
        "2.2. Critères éliminatoires",
        "III. Résultats de la présélection",
    ]
    assert len(rendues) == len(entrees)


def test_une_section_vide_ne_consomme_pas_de_numero():
    """Sinon le document sauterait du I au III sans que rien ne l'explique."""
    sections = trame()
    # Ni texte, ni tableau, ni sous-sections à porter : elle ne paraît pas.
    sections[2]["porteur"] = False
    del sections[3:5]

    _, entrees, _ = export.preparer("T", "C", sections, QUAND, GARDE)
    assert [e.titre_complet for e in entrees] == [
        "Introduction",
        "I. Démarche",
        "II. Résultats de la présélection",
    ]


def test_une_sous_section_sans_section_porteuse_se_numerote_a_plat():
    """Un modèle client peut ouvrir sur une sous-section. Pas de « 0.1 »."""
    _, entrees, _ = export.preparer(
        "T", "C", [section("A", "Seule", niveau=2)], QUAND, GARDE
    )
    assert entrees[0].titre_complet == "1. Seule"


def test_le_titre_ne_repete_pas_l_objet_deja_ecrit_au_dessus():
    """« Rapport final — Cadres 2026 » sous « CADRES 2026 » fait bafouiller.

    Le titre complet reste celui de la liste des rapports, où le nom du mandat
    est la seule chose qui distingue deux lignes.
    """
    garde = frontispice.Couverture(
        titre="Rapport final de recrutement — Cadres 2026", objet="Cadres 2026"
    )
    assert garde.titre_document == "Rapport final de recrutement"
    assert garde.titre == "Rapport final de recrutement — Cadres 2026"

    # Un titre qui ne reprend pas l'objet n'est pas touché.
    autre = frontispice.Couverture(titre="Rapport des entretiens", objet="Cadres 2026")
    assert autre.titre_document == "Rapport des entretiens"

    # Et un titre qui n'est *que* l'objet garde quelque chose à afficher.
    seul = frontispice.Couverture(titre="Cadres 2026", objet="Cadres 2026")
    assert seul.titre_document == "Cadres 2026"


def test_le_mois_est_celui_de_la_page_de_garde_du_cabinet():
    assert frontispice.mois_de(datetime(2025, 5, 12)) == "Mai 2025"
    assert frontispice.mois_de(datetime(2026, 8, 1)) == "Août 2026"


# --- la page de garde, dans les quatre formats ------------------------------


def test_le_docx_porte_la_page_de_garde_sur_le_papier_a_en_tete():
    """L'objet, le titre et le mois se posent SUR l'en-tête du cabinet.

    Les coordonnées ne sont plus composées en texte : elles font partie de
    l'image d'en-tête, comme dans le document que le cabinet remet. Les
    réécrire les ferait paraître deux fois.
    """
    octets = rendre(export.rendre_docx)
    document = Document(io.BytesIO(octets))
    textes = [p.text for p in document.paragraphs]

    assert GARDE.objet_affiche in textes
    assert GARDE.titre in textes
    assert "Réf. WAPP/2025/DRH-07" in textes
    assert "Mai 2025" in textes

    # Le papier à en-tête est posé en fond de première page, ancré à la page
    # et derrière le texte : en ligne, il resterait au ras du pied.
    #
    # Le numéro de la part n'est pas fixé — Word en range deux, le pied
    # courant et celui de première page — donc on les parcourt toutes plutôt
    # que de parier sur « footer1 ».
    with zipfile.ZipFile(io.BytesIO(octets)) as archive:
        pieds = [
            archive.read(n).decode()
            for n in archive.namelist()
            if n.startswith("word/footer") and n.endswith(".xml")
        ]
    flottants = [
        p for p in pieds if "<wp:anchor" in p and 'behindDoc="1"' in p
    ]
    assert flottants, "l'en-tête doit flotter derrière le texte, pas être en ligne"
    assert 'relativeFrom="page"' in flottants[0]


def test_le_docx_dresse_un_vrai_champ_de_table_des_matieres():
    """Word et LibreOffice pagineront eux-mêmes ; le cache tient l'attente."""
    octets = rendre(export.rendre_docx)
    with zipfile.ZipFile(io.BytesIO(octets)) as archive:
        corps = archive.read("word/document.xml").decode()
        reglages = archive.read("word/settings.xml").decode()

    assert 'TOC \\o "1-2"' in corps
    assert "<w:updateFields" in reglages, "le sommaire doit s'actualiser à l'ouverture"

    # Et les titres paraissent dès l'ouverture, même sans actualisation.
    textes = [p.text for p in Document(io.BytesIO(octets)).paragraphs]
    assert export.TITRE_SOMMAIRE in textes
    assert "I. Démarche" in textes


def test_le_docx_numerote_ses_pages_hors_page_de_garde():
    document = Document(io.BytesIO(rendre(export.rendre_docx)))
    assert document.sections[0].different_first_page_header_footer
    # La section du corps ne doit pas hériter du pied de garde, sinon le
    # sommaire reprend l'adresse du cabinet au lieu de sa pagination.
    assert not document.sections[1].different_first_page_header_footer
    assert frontispice.MENTION_CONFIDENTIEL in document.sections[0].footer.paragraphs[0].text


def test_les_titres_du_docx_sont_des_titres_de_niveau_1_et_2():
    """C'est ce que le champ TOC recense. Un cran de décalage le laissait vide."""
    document = Document(io.BytesIO(rendre(export.rendre_docx)))
    par_texte = {p.text: p.style.name for p in document.paragraphs}
    assert par_texte["I. Démarche"] == "Heading 1"
    assert par_texte["2.1. Présélection"] == "Heading 2"


def test_l_odt_porte_la_page_de_garde_le_logo_et_un_index():
    octets = rendre(export.rendre_odt)
    with zipfile.ZipFile(io.BytesIO(octets)) as archive:
        contenu = archive.read("content.xml").decode()
        styles = archive.read("styles.xml").decode()
        logo = archive.read("Pictures/logo_kapi.png")
        manifeste = archive.read("META-INF/manifest.xml").decode()

    assert GARDE.objet_affiche in contenu
    assert frontispice.ADRESSE[0] in contenu

    # Un vrai index ODF : LibreOffice en refait la pagination à la demande, et
    # son corps porte les titres dès l'ouverture.
    assert "<text:table-of-content " in contenu
    assert "text:index-entry-page-number" in contenu
    assert "I. Démarche" in contenu

    # Le logo est embarqué et déclaré au manifeste, sans quoi l'ODT est refusé.
    assert logo.startswith(b"\x89PNG\r\n\x1a\n")
    assert "Pictures/logo_kapi.png" in manifeste
    # Le bandeau du cabinet court au pied de chaque page, comme dans le
    # document remis, et il est déclaré lui aussi.
    assert "Pictures/bandeau_kapi.png" in manifeste
    assert "Pictures/bandeau_kapi.png" in styles
    # Le pied de page vit dans la page maîtresse.
    assert "text:page-number" in styles


def test_le_pdf_pagine_son_sommaire():
    """Le seul format où la pagination est calculée par nous, en deux passes."""
    octets = rendre(export.rendre_pdf)
    assert octets.startswith(b"%PDF")
    # Page de garde, sommaire, puis le corps : trois pages au moins.
    assert octets.count(b"/Type /Page\n") >= 3 or octets.count(b"/Type/Page") >= 3


def test_le_txt_porte_la_garde_et_la_liste_des_titres():
    texte = rendre(export.rendre_txt).decode("utf-8")
    assert frontispice.NOM_CABINET in texte
    assert GARDE.objet_affiche in texte
    assert frontispice.ADRESSE[0] in texte
    assert export.TITRE_SOMMAIRE.upper() in texte
    # Le titre paraît deux fois : au sommaire, puis en tête de sa section.
    assert texte.count("I. Démarche") == 1
    assert "I. DÉMARCHE" in texte


@pytest.mark.parametrize(
    "rendu", [export.rendre_docx, export.rendre_pdf, export.rendre_odt, export.rendre_txt]
)
def test_un_appel_sans_couverture_en_dresse_une_quand_meme(rendu):
    """Un script ou un test n'a que le titre : il doit sortir un vrai document."""
    octets = rendu("Rapport", "WAPP", trame(), QUAND)
    assert octets
    if rendu is export.rendre_txt:
        assert "Mai 2025" in octets.decode("utf-8")


@pytest.mark.parametrize(
    "rendu", [export.rendre_docx, export.rendre_pdf, export.rendre_odt, export.rendre_txt]
)
def test_un_rapport_sans_aucune_section_sort_sa_page_de_garde(rendu):
    """Un brouillon vierge s'exporte : la garde ne dépend pas de la rédaction."""
    assert rendu("Rapport", "WAPP", [], QUAND, GARDE)


# --- le DOCX est conforme au schéma OOXML ------------------------------------
#
# `w:pPr`, `w:rPr` et `w:tcPr` ne sont pas des sacs d'attributs : le schéma
# OOXML en décrit la *séquence*. `w:pBdr` — le filet or sous la marque — était
# ajouté à la fin des propriétés du paragraphe, donc après le `w:spacing` et le
# `w:jc` posés juste avant. `w:updateFields` avait le même défaut.
#
# Ce que ce contrôle ne prétend PAS : que Word refuserait le document. On l'a
# cru, puis mesuré — Word 16 ouvre les deux ordres sans broncher, réparation
# automatique désactivée. La conformité se tient pour elle-même : rien ne
# garantit la même indulgence chez un validateur, une bibliothèque tierce ou
# une version future, et un document conforme ne coûte rien de plus.

_ORDRE_PPR = """pStyle keepNext keepLines pageBreakBefore framePr widowControl
numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap
overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd
snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc
textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr
sectPr pPrChange""".split()

_ORDRE_RPR = """rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline
shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern
position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang
eastAsianLayout specVanish oMath rPrChange""".split()

_ORDRE_TCPR = """cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar
textDirection tcFitText vAlign hideMark headers cellIns cellDel cellMerge
tcPrChange""".split()

# `w:updateFields`, posé pour que le sommaire se recalcule, a la même
# contrainte : il vient avant `w:compat` dans `CT_Settings`.
_ORDRE_SETTINGS = """zoom proofState defaultTabStop characterSpacingControl
savePreviewPicture updateFields compat rsids mathPr themeFontLang
clrSchemeMapping doNotAutoCompressPictures shapeDefaults decimalSymbol
listSeparator""".split()


def _desordres(xml: str, conteneur: str, ordre: list[str]) -> list[str]:
    """Les enfants de `conteneur` qui rompent la séquence du schéma."""
    fautes = []
    motif = rf"<w:{conteneur}(?:\s[^>]*)?>(.*?)</w:{conteneur}>"
    for bloc in re.finditer(motif, xml, re.S):
        vus = [
            nom
            for nom in re.findall(r"<w:([a-zA-Z]+)[\s/>]", bloc.group(1))
            if nom in ordre
        ]
        rangs = [ordre.index(nom) for nom in vus]
        fautes += [
            f"{conteneur} : {vus[i]} après {vus[i - 1]}"
            for i in range(1, len(rangs))
            if rangs[i] < rangs[i - 1]
        ]
    return fautes


def test_le_docx_respecte_la_sequence_du_schema_ooxml():
    """Conformité au schéma — et rien d'autre dans la suite ne la contrôle."""
    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_docx))) as archive:
        corps = archive.read("word/document.xml").decode()
        reglages = archive.read("word/settings.xml").decode()

    fautes = (
        _desordres(corps, "pPr", _ORDRE_PPR)
        + _desordres(corps, "rPr", _ORDRE_RPR)
        + _desordres(corps, "tcPr", _ORDRE_TCPR)
        + _desordres(reglages, "settings", _ORDRE_SETTINGS)
    )
    assert not fautes, "éléments hors séquence OOXML : " + " ; ".join(fautes)


def test_le_docx_se_relit_et_porte_tout_le_texte():
    """Une corruption d'ordre ne se voit pas ici : ce test ne suffit pas seul."""
    sections = trame()
    document = Document(io.BytesIO(rendre(export.rendre_docx, sections)))
    rendu = "\n".join(p.text for p in document.paragraphs)
    for s in sections:
        assert s["contenu"] in rendu or not s["contenu"]


def test_l_odt_range_ses_proprietes_dans_l_ordre_du_schema_odf():
    """`style:paragraph-properties` avant `style:text-properties`.

    ODF décrit une séquence, et l'ordre inverse coûte cher parce qu'il ne casse
    rien de visible : un lecteur qui applique le schéma jette ce qui tient au
    paragraphe et garde ce qui tient au texte. La page de garde sortait donc
    ferrée à gauche et tassée en haut, avec la bonne police et les bonnes
    couleurs — aucun style ne manquait à l'appel, ils étaient tous à moitié
    appliqués.
    """
    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_odt))) as archive:
        styles = archive.read("styles.xml").decode()

    fautifs = []
    for bloc in re.finditer(r"<style:style\b.*?</style:style>", styles, re.S):
        texte = bloc.group(0)
        nom = re.search(r'style:name="([^"]+)"', texte)
        pos_p = texte.find("<style:paragraph-properties")
        pos_t = texte.find("<style:text-properties")
        if pos_p != -1 and pos_t != -1 and pos_p > pos_t:
            fautifs.append(nom.group(1) if nom else "?")

    assert not fautifs, (
        "styles ODF dont les propriétés de paragraphe suivent celles de texte : "
        + ", ".join(fautifs)
    )


def test_l_odt_pose_sa_police_et_son_interligne():
    """Sans quoi il sort dans la serif par défaut du lecteur, en bloc compact.

    Le DOCX est en Calibri avec 8 points entre paragraphes ; l'ODT n'imposait
    rien et LibreOffice le rendait en Liberation Serif, paragraphes collés.
    Trois formats du même rapport n'ont pas à se ressembler de loin seulement.
    """
    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_odt))) as archive:
        styles = archive.read("styles.xml").decode()

    assert "<office:font-face-decls>" in styles
    # Carlito prend le relais là où Calibri n'est pas installé, sans changer
    # la mise en page : c'est son clone métrique.
    assert "Carlito" in styles

    defaut = re.search(
        r"<style:default-style style:family=\"paragraph\">.*?</style:default-style>",
        styles,
        re.S,
    )
    assert defaut, "aucun style de paragraphe par défaut"
    assert "fo:margin-bottom" in defaut.group(0), "pas d'interligne entre paragraphes"
    assert 'style:font-name="Calibri"' in defaut.group(0)


def test_la_page_de_garde_odt_est_bien_centree():
    """Le centrage est ce que l'inversion faisait perdre en premier."""
    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_odt))) as archive:
        styles = archive.read("styles.xml").decode()

    for nom in ("GardeMarque", "GardeObjet", "GardeTitre", "GardeClient", "GardeMois"):
        bloc = re.search(
            rf'<style:style style:name="{nom}".*?</style:style>', styles, re.S
        )
        assert bloc, f"style {nom} absent"
        assert 'fo:text-align="center"' in bloc.group(0), f"{nom} n'est pas centré"


# --- les énumérations --------------------------------------------------------


def test_une_puce_devient_une_vraie_liste_et_perd_son_tiret():
    """Le tiret rendu tel quel donnait des paragraphes ordinaires, sans retrait.

    La deuxième ligne d'un élément long revenait alors à la marge et
    l'énumération se perdait à la lecture.
    """
    section = {
        "code": "OBJECTIFS",
        "titre": "Objectifs",
        "contenu": "Le but général est de :\n- analyser les dossiers ;\n- assister le jury.",
        "niveau": 1,
    }
    assert export.paragraphes_de(section) == [
        ("Le but général est de :", False),
        ("analyser les dossiers ;", True),
        ("assister le jury.", True),
    ]

    document = Document(io.BytesIO(rendre(export.rendre_docx, [section])))
    puces = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert [p.text for p in puces] == ["analyser les dossiers ;", "assister le jury."]

    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_odt, [section]))) as archive:
        contenu = archive.read("content.xml").decode()
    assert "<text:list text:style-name=\"Puces\">" in contenu
    assert "- analyser" not in contenu


@pytest.mark.parametrize("marque", ["- ", "– ", "— ", "• ", "* "])
def test_les_marques_d_enumeration_usuelles_sont_reconnues(marque):
    """Le modèle n'emploie pas toujours le même signe d'une section à l'autre."""
    section = {"code": "X", "titre": "X", "contenu": f"{marque}un élément"}
    assert export.paragraphes_de(section) == [("un élément", True)]


# --- le commentaire sous le tableau ------------------------------------------
#
# Le document du cabinet ne met pas ses tableaux en fin de section : une phrase
# les annonce, le tableau suit, un commentaire des chiffres vient après —
# « Trente-trois (33) candidatures préqualifiées pour le poste de DAF, dont
# cinq (5) proposés pour la prochaine étape ». Tout rendre avant le tableau
# plaçait ce commentaire au-dessus des chiffres qu'il commente, juste après la
# phrase annonçant « les résultats suivants ».


def section_a_tableau() -> dict:
    return {
        "code": "RESULTATS_PRESELECTION",
        "titre": "Résultats de la présélection",
        "contenu": "L'analyse a permis d'obtenir les résultats suivants :",
        "contenu_apres": "Trente-trois (33) candidatures ont été préqualifiées.",
        "niveau": 1,
        "tableaux": [
            {
                "titre": "Effectifs",
                "colonnes": [{"libelle": "Postes"}, {"libelle": "Reçus", "numerique": True}],
                "lignes": [{"cellules": ["DAF", "64"], "genre": ""}],
            }
        ],
    }


def test_le_commentaire_se_lit_apres_le_tableau_dans_le_docx():
    document = Document(io.BytesIO(rendre(export.rendre_docx, [section_a_tableau()])))

    # L'ordre des blocs du corps : annonce, table, commentaire.
    ordre = []
    for enfant in document.element.body.iterchildren():
        if enfant.tag.endswith("}tbl"):
            ordre.append("TABLE")
        elif enfant.tag.endswith("}p"):
            texte = "".join(enfant.itertext()).strip()
            if texte.startswith("L'analyse"):
                ordre.append("AVANT")
            elif texte.startswith("Trente-trois"):
                ordre.append("APRES")
    assert ordre == ["AVANT", "TABLE", "APRES"]


def test_le_commentaire_se_lit_apres_le_tableau_dans_l_odt_et_le_txt():
    with zipfile.ZipFile(
        io.BytesIO(rendre(export.rendre_odt, [section_a_tableau()]))
    ) as archive:
        contenu = archive.read("content.xml").decode()
    assert contenu.index("L'analyse") < contenu.index("<table:table ")
    assert contenu.index("<table:table ") < contenu.index("Trente-trois")

    texte = rendre(export.rendre_txt, [section_a_tableau()]).decode("utf-8")
    # Le TXT porte le tableau aligné à l'espace ; le commentaire vient après.
    assert texte.index("L'analyse") < texte.index("Trente-trois")


def test_une_section_qui_n_a_que_du_texte_sous_le_tableau_parait_quand_meme():
    """Sinon la section disparaîtrait, emportant son tableau et son commentaire."""
    section = dict(section_a_tableau(), contenu="")
    _, entrees, _ = export.preparer("T", "C", [section], QUAND, GARDE)
    assert len(entrees) == 1


# --- les pièces de marque du cabinet -----------------------------------------
#
# Elles ne sont pas redessinées : elles sont extraites des documents que le
# cabinet remet. Le logo l'avait d'abord été — neuf carrés en dégradé, tracés
# pixel par pixel d'après le site — quand le vrai en compte seize, en damier,
# et porte le mot « Kapi Consult » à côté. Le rapport sortait donc sous une
# marque que le cabinet n'emploie pas.


def test_les_trois_pieces_de_marque_accompagnent_le_code():
    assert frontispice.pieces_presentes()
    assert frontispice.logo_png().startswith(b"\x89PNG\r\n\x1a\n")
    assert frontispice.bandeau_png().startswith(b"\x89PNG\r\n\x1a\n")
    assert frontispice.entete_jpg().startswith(b"\xff\xd8\xff")


def test_le_pdf_pose_le_papier_a_en_tete_en_fond():
    """La première page est celle du cabinet, pas une imitation.

    Un JPEG embarqué se reconnaît à son filtre de décodage : `DCTDecode`
    n'apparaît que si le PDF porte une image JPEG.
    """
    assert b"DCTDecode" in rendre(export.rendre_pdf)


def test_le_docx_embarque_l_en_tete_et_le_bandeau():
    with zipfile.ZipFile(io.BytesIO(rendre(export.rendre_docx))) as archive:
        medias = [n for n in archive.namelist() if n.startswith("word/media/")]
    assert len(medias) >= 2, "le papier à en-tête et le bandeau du pied"
