from __future__ import annotations

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditLog

pytestmark = pytest.mark.anyio

API = "/api/v1"


async def monter_poste(client, auth, **overrides) -> str:
    reponse = await client.post(
        f"{API}/clients", json={"nom": "Dogta-Lafiè", "secteur": "Santé"}, headers=auth
    )
    assert reponse.status_code == 201, reponse.text
    client_id = reponse.json()["id"]

    reponse = await client.post(
        f"{API}/mandats",
        json={"client_id": client_id, "intitule": "Direction générale"},
        headers=auth,
    )
    assert reponse.status_code == 201, reponse.text
    mandat_id = reponse.json()["id"]

    corps = {
        "intitule": "Directeur Général",
        "niveau_min": 4,
        "domaines_acceptes": ["gestion hôtelière"],
        "annees_experience_min": 10,
        "annees_experience_specifique_min": 5,
        "domaines_experience": ["gestion hôtelière"],
        "pieces_requises": ["LETTRE_MOTIVATION", "CV"],
        "langues_requises": ["français"],
    }
    corps.update(overrides)
    reponse = await client.post(f"{API}/mandats/{mandat_id}/postes", json=corps, headers=auth)
    assert reponse.status_code == 201, reponse.text
    poste_id = reponse.json()["id"]

    reponse = await client.post(
        f"{API}/postes/{poste_id}/avis",
        json={"type_avis": "NATIONAL", "date_cloture": "2026-07-31"},
        headers=auth,
    )
    assert reponse.status_code == 201, reponse.text
    return poste_id


def dossier(**overrides) -> dict:
    candidat = {
        "nom": "Kodjo",
        "prenom": "Amina",
        # Exigée à la saisie : c'est par là qu'on accuse réception, qu'on
        # réclame une pièce et qu'on convoque, et elle figure en colonne dans
        # le tableau des préqualifiés remis au client.
        "email": "amina.kodjo@example.tg",
        "date_naissance": "1985-03-01",
        "sexe": "F",
        "nationalites": ["Togolaise"],
        "adresse": "Lomé, Tokoin",
        "langues": ["français"],
        "diplomes": [
            {
                "intitule": "Master en gestion hôtelière",
                "niveau": 5,
                "domaine": "gestion hôtelière",
                "etablissement": "Université de Lomé",
                "annee": 2010,
            }
        ],
        "experiences": [
            {
                "poste": "Directrice adjointe",
                "employeur": "Hôtel du 2 Février",
                "debut": "2010-01-01",
                "fin": "2026-01-01",
                "domaines": ["gestion hôtelière"],
                "pays": "Togo",
            }
        ],
    }
    pieces = overrides.pop("pieces_fournies", ["LETTRE_MOTIVATION", "CV"])
    candidat.update(overrides)
    return {
        "candidat": candidat,
        "recue_le": "2026-07-01T10:00:00",
        "pieces_fournies": pieces,
    }


async def test_chaine_complete_jusqu_a_la_grille(client, auth):
    poste_id = await monter_poste(client, auth)

    reponse = await client.post(
        f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth
    )
    assert reponse.status_code == 201, reponse.text
    corps = reponse.json()
    assert corps["statut"] == "PRESELECTIONNEE"
    assert corps["notation"]["total"] > 0
    assert len(corps["notation"]["lignes"]) == 4

    grille = await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)
    assert grille.status_code == 200, grille.text
    donnees = grille.json()
    assert donnees["nombre_candidatures"] == 1
    assert donnees["nombre_preselectionnes"] == 1
    ligne = donnees["preselectionnes"][0]
    # Les colonnes du fichier de travail.
    assert ligne["nom"] == "Kodjo"
    assert ligne["age"] == 41
    assert ligne["nationalite"] == "Togolaise"
    assert ligne["dernier_diplome"] == "Master en gestion hôtelière"
    assert ligne["ecole_universite"] == "Université de Lomé"
    assert ligne["structure_employeur"] == "Hôtel du 2 Février"


async def test_le_dossier_incomplet_apparait_dans_le_tableau_d_elimination(client, auth):
    poste_id = await monter_poste(client, auth)
    await client.post(
        f"{API}/postes/{poste_id}/candidatures",
        json=dossier(pieces_fournies=["CV"]),  # lettre de motivation manquante
        headers=auth,
    )

    donnees = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    groupes = {g["motif"]: g for g in donnees["elimines"]}
    assert "DOSSIER_INCOMPLET" in groupes
    assert groupes["DOSSIER_INCOMPLET"]["libelle"] == "Dossier incomplet"


async def test_lever_un_motif_reclasse_la_candidature(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(
            f"{API}/postes/{poste_id}/candidatures",
            json=dossier(pieces_fournies=["CV"]),
            headers=auth,
        )
    ).json()["id"]

    reponse = await client.post(
        f"{API}/candidatures/{candidature_id}/eliminations/DOSSIER_INCOMPLET/lever",
        json={"motif": "Pièces reçues séparément par email."},
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["leve_le"] is not None

    detail = (await client.get(f"{API}/candidatures/{candidature_id}", headers=auth)).json()
    assert detail["statut"] == "PRESELECTIONNEE"

    # Une seconde levée est refusée plutôt qu'appliquée deux fois.
    rejouee = await client.post(
        f"{API}/candidatures/{candidature_id}/eliminations/DOSSIER_INCOMPLET/lever",
        json={"motif": "encore"},
        headers=auth,
    )
    assert rejouee.status_code == 409


async def test_une_condition_restrictive_sans_justification_est_refusee(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.patch(
        f"{API}/postes/{poste_id}",
        json={"restriction": {"age_max": 35, "justification": ""}},
        headers=auth,
    )
    assert reponse.status_code == 422


async def test_une_condition_restrictive_justifiee_est_journalisee(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.patch(
        f"{API}/postes/{poste_id}",
        json={
            "restriction": {
                "age_max": 35,
                "justification": "Limite d'âge du statut du personnel.",
            }
        },
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["restriction"]["age_max"] == 35

    await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    donnees = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    motifs = {g["motif"] for g in donnees["elimines"]}
    assert "CONDITION_AGE" in motifs


async def test_abaisser_le_seuil_exige_une_justification(client, auth):
    # Le plancher vaut zéro par défaut — le classement sélectionne. Il faut
    # donc qu'un plancher existe pour qu'on puisse l'abaisser : on le pose sur
    # le poste, ce qui est désormais le seul endroit où il se règle.
    poste_id = await monter_poste(client, auth)
    pose = await client.post(
        f"{API}/postes/{poste_id}/seuil", json={"seuil": 20}, headers=auth
    )
    assert pose.status_code == 200, pose.text

    refus = await client.post(
        f"{API}/postes/{poste_id}/seuil", json={"seuil": 15, "justification": ""}, headers=auth
    )
    assert refus.status_code == 422

    accepte = await client.post(
        f"{API}/postes/{poste_id}/seuil",
        json={"seuil": 15, "justification": "Poste en tension, trois candidatures."},
        headers=auth,
    )
    assert accepte.status_code == 200
    assert accepte.json()["seuil_preselection"] == 15.0


async def test_un_bareme_incoherent_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.put(
        f"{API}/postes/{poste_id}/bareme",
        json={
            "bareme": {
                "formation": {"points_max": 10, "points_niveau_requis": 8},
                "experience_generale": {"points_max": 5, "points_au_seuil": 3},
                "experience_specifique": {"points_max": 5, "points_au_seuil": 3},
                "points_ajouts_max": 0,
                "total_max": 30,
            }
        },
        headers=auth,
    )
    assert reponse.status_code == 422
    assert "plafonne" in reponse.json()["detail"]


async def test_experience_specifique_superieure_a_la_generale_est_refusee(client, auth):
    """Une incohérence qui rendrait le poste impossible à satisfaire."""
    reponse = await client.post(
        f"{API}/clients", json={"nom": "X"}, headers=auth
    )
    client_id = reponse.json()["id"]
    mandat_id = (
        await client.post(
            f"{API}/mandats", json={"client_id": client_id, "intitule": "M"}, headers=auth
        )
    ).json()["id"]

    reponse = await client.post(
        f"{API}/mandats/{mandat_id}/postes",
        json={
            "intitule": "Poste",
            "annees_experience_min": 3,
            "annees_experience_specifique_min": 8,
        },
        headers=auth,
    )
    assert reponse.status_code == 422


async def test_publier_un_avis_sans_cloture_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    avis_id = (
        await client.post(
            f"{API}/postes/{poste_id}/avis", json={"type_avis": "INTERNATIONAL"}, headers=auth
        )
    ).json()["id"]

    refus = await client.post(f"{API}/avis/{avis_id}/publier", headers=auth)
    assert refus.status_code == 422
    assert "clôture" in refus.json()["detail"]

    await client.patch(f"{API}/avis/{avis_id}", json={"date_cloture": "2026-09-30"}, headers=auth)
    publie = await client.post(f"{API}/avis/{avis_id}/publier", headers=auth)
    assert publie.status_code == 200
    assert publie.json()["statut"] == "PUBLIE"
    assert publie.json()["date_publication"] is not None


async def test_liste_filtrable_et_paginee(client, auth):
    poste_id = await monter_poste(client, auth)
    for prenom in ("Amina", "Bernard", "Chantal"):
        await client.post(
            f"{API}/postes/{poste_id}/candidatures", json=dossier(prenom=prenom), headers=auth
        )

    page = (
        await client.get(f"{API}/postes/{poste_id}/candidatures", headers=auth)
    ).json()
    assert page["total"] == 3
    assert len(page["items"]) == 3
    assert page["items"][0]["note"] is not None
    # Dossiers complets et conformes : aucun motif retenu.
    assert page["items"][0]["motifs"] == []
    assert page["items"][0]["statut"] == "PRESELECTIONNEE"

    recherche = (
        await client.get(
            f"{API}/postes/{poste_id}/candidatures?recherche=Bernard", headers=auth
        )
    ).json()
    assert recherche["total"] == 1

    petite = (
        await client.get(f"{API}/postes/{poste_id}/candidatures?page_size=2", headers=auth)
    ).json()
    assert len(petite["items"]) == 2
    assert petite["total"] == 3


async def test_note_manuelle_exige_un_motif_et_prime(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    ).json()["id"]

    refus = await client.patch(
        f"{API}/candidatures/{candidature_id}/note", json={"note": 27}, headers=auth
    )
    assert refus.status_code == 422

    accepte = await client.patch(
        f"{API}/candidatures/{candidature_id}/note",
        json={"note": 27, "motif": "Entretien téléphonique concluant."},
        headers=auth,
    )
    assert accepte.status_code == 200
    assert accepte.json()["notation"]["note_manuelle"] == 27.0

    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    lignes = grille["preselectionnes"] + grille["non_retenus"]
    lignes += [l for g in grille["elimines"] for l in g["lignes"]]
    assert lignes[0]["note"] == 27.0


async def test_reevaluer_le_poste_est_idempotent(client, auth):
    poste_id = await monter_poste(client, auth)
    await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)

    for _ in range(3):
        reponse = await client.post(f"{API}/postes/{poste_id}/evaluer", headers=auth)
        assert reponse.status_code == 200
        assert reponse.json()["candidatures_evaluees"] == 1

    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    assert grille["nombre_candidatures"] == 1
    motifs = [g["motif"] for g in grille["elimines"]]
    assert len(motifs) == len(set(motifs))


PDF = b"%PDF-1.4\n" + b"0" * 400


async def test_joindre_un_fichier_complete_le_dossier(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(
            f"{API}/postes/{poste_id}/candidatures",
            json=dossier(pieces_fournies=["CV"]),
            headers=auth,
        )
    ).json()["id"]

    reponse = await client.post(
        f"{API}/candidatures/{candidature_id}/pieces",
        data={"type_piece": "LETTRE_MOTIVATION"},
        files={"fichier": ("lettre.pdf", PDF, "application/pdf")},
        headers=auth,
    )
    assert reponse.status_code == 201, reponse.text
    corps = reponse.json()
    # Le dossier devient complet : le motif tombe de lui-même.
    assert corps["statut"] == "PRESELECTIONNEE"
    assert {p["type_piece"] for p in corps["pieces"]} == {"CV", "LETTRE_MOTIVATION"}


async def test_le_fichier_se_rattache_a_la_piece_deja_constatee(client, auth):
    """Cocher « CV reçu » puis déposer le fichier ne crée pas deux pièces."""
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(
            f"{API}/postes/{poste_id}/candidatures",
            json=dossier(pieces_fournies=["CV", "LETTRE_MOTIVATION"]),
            headers=auth,
        )
    ).json()["id"]

    corps = (
        await client.post(
            f"{API}/candidatures/{candidature_id}/pieces",
            data={"type_piece": "CV"},
            files={"fichier": ("cv.pdf", PDF, "application/pdf")},
            headers=auth,
        )
    ).json()
    assert len([p for p in corps["pieces"] if p["type_piece"] == "CV"]) == 1
    piece = next(p for p in corps["pieces"] if p["type_piece"] == "CV")
    assert piece["nom_fichier"] == "cv.pdf"


async def test_telecharger_puis_retirer_une_piece(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    ).json()["id"]

    corps = (
        await client.post(
            f"{API}/candidatures/{candidature_id}/pieces",
            data={"type_piece": "CV"},
            files={"fichier": ("cv.pdf", PDF, "application/pdf")},
            headers=auth,
        )
    ).json()
    piece_id = next(p["id"] for p in corps["pieces"] if p["type_piece"] == "CV")

    fichier = await client.get(
        f"{API}/candidatures/{candidature_id}/pieces/{piece_id}", headers=auth
    )
    assert fichier.status_code == 200
    assert fichier.content == PDF

    retrait = await client.request(
        "DELETE", f"{API}/candidatures/{candidature_id}/pieces/{piece_id}", headers=auth
    )
    assert retrait.status_code == 200
    assert all(p["id"] != piece_id for p in retrait.json()["pieces"])


async def test_supprimer_un_dossier_entre_par_erreur(client, auth):
    """Le cas courant : un dossier saisi en double, qui fausse les comptes."""
    poste_id = await monter_poste(client, auth)
    # Incomplet, donc éliminé : rien ne retient sa suppression.
    candidature_id = (
        await client.post(
            f"{API}/postes/{poste_id}/candidatures",
            json=dossier(pieces_fournies=["CV"]),
            headers=auth,
        )
    ).json()["id"]
    await client.post(
        f"{API}/candidatures/{candidature_id}/pieces",
        data={"type_piece": "CV"},
        files={"fichier": ("cv.pdf", PDF, "application/pdf")},
        headers=auth,
    )

    reponse = await client.request(
        "DELETE",
        f"{API}/candidatures/{candidature_id}",
        params={"motif": "saisi deux fois, celui-ci est le doublon"},
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    corps = reponse.json()
    assert corps["pieces_supprimees"] == 1
    # Le profil n'existait que pour ce dossier : il part avec lui, plutôt que
    # de rester dans le vivier sans pièce ni candidature.
    assert corps["profil_supprime"] is True

    assert (
        await client.get(f"{API}/candidatures/{candidature_id}", headers=auth)
    ).status_code == 404
    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    assert grille["nombre_candidatures"] == 0

    # Le dossier a disparu, pas la trace de sa disparition.
    async with SessionLocal() as db:
        entree = (
            await db.execute(
                select(AuditLog).where(AuditLog.action == "candidature.supprimer")
            )
        ).scalar_one()
        assert entree.entity_id == candidature_id
        assert entree.details["motif"] == "saisi deux fois, celui-ci est le doublon"
        assert entree.details["candidat"] == "Kodjo Amina"
        assert entree.details["pieces"] == ["cv.pdf"]


async def test_supprimer_un_dossier_sans_motif_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    ).json()["id"]

    reponse = await client.request(
        "DELETE", f"{API}/candidatures/{candidature_id}", headers=auth
    )
    assert reponse.status_code == 422


async def test_un_dossier_preselectionne_resiste_puis_cede(client, auth):
    """Ce qui a déjà servi ne s'efface pas d'un clic, mais reste effaçable.

    Le refus n'est pas un verrou : il oblige à voir ce que le dossier porte
    avant de le confirmer.
    """
    poste_id = await monter_poste(client, auth)
    depot = (
        await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    ).json()
    assert depot["statut"] == "PRESELECTIONNEE"

    refus = await client.request(
        "DELETE",
        f"{API}/candidatures/{depot['id']}",
        params={"motif": "erreur de poste"},
        headers=auth,
    )
    assert refus.status_code == 409
    assert "présélection" in refus.json()["detail"]

    accepte = await client.request(
        "DELETE",
        f"{API}/candidatures/{depot['id']}",
        params={"motif": "erreur de poste", "confirmer": "true"},
        headers=auth,
    )
    assert accepte.status_code == 200, accepte.text

    async with SessionLocal() as db:
        entree = (
            await db.execute(
                select(AuditLog).where(AuditLog.action == "candidature.supprimer")
            )
        ).scalar_one()
        # Le journal dit ce sur quoi on est passé outre.
        assert entree.details["passe_outre"] == ["il figure dans la présélection"]


async def test_un_fichier_non_pdf_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    candidature_id = (
        await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    ).json()["id"]

    reponse = await client.post(
        f"{API}/candidatures/{candidature_id}/pieces",
        data={"type_piece": "CV"},
        files={"fichier": ("photo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 200, "image/png")},
        headers=auth,
    )
    assert reponse.status_code == 422


async def test_etat_du_courriel_sans_boite_configuree(client, auth):
    etat = await client.get(f"{API}/courriel/etat", headers=auth)
    assert etat.status_code == 200
    assert etat.json()["actif"] is False

    releve = await client.post(f"{API}/courriel/relever", headers=auth)
    assert releve.status_code == 409


async def test_export_excel_de_la_grille(client, auth):
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    await client.post(f"{API}/postes/{poste_id}/candidatures", json=dossier(), headers=auth)
    await client.post(
        f"{API}/postes/{poste_id}/candidatures",
        json=dossier(nom="Abalo", prenom="Yao", pieces_fournies=["CV"]),
        headers=auth,
    )

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    assert reponse.status_code == 200, reponse.text
    assert "spreadsheetml" in reponse.headers["content-type"]
    # « dossier-complet » et non « grille-preselection » : depuis que la grille
    # de présélection s'exporte seule, deux fichiers différents porteraient
    # sinon le même nom dans le dossier de téléchargement.
    assert "dossier-complet-" in reponse.headers["content-disposition"]

    classeur = load_workbook(io.BytesIO(reponse.content))
    assert classeur.sheetnames == [
        "Grille de présélection",
        "Tableau d'élimination",
        "Synthèse",
    ]

    grille = classeur["Grille de présélection"]
    entete = {row[0] for row in grille.iter_rows(min_col=1, max_col=1, values_only=True)}
    # L'en-tête doit permettre de relire la grille des mois plus tard.
    assert "Client" in entete
    assert "Date de clôture" in entete
    # Le barème rappelle que la présélection ne pèse que 30 % de la note
    # finale : une note sur 30 lue seule se prendrait pour un résultat.
    assert "Barème de présélection" in entete
    assert any(cell and "Classement des dossiers" in str(cell) for cell in entete)

    contenu = [
        c
        for row in grille.iter_rows(values_only=True)
        for c in row
        if isinstance(c, str)
    ]
    assert "Kodjo" in contenu

    elimination = classeur["Tableau d'élimination"]
    motifs = [
        c
        for row in elimination.iter_rows(min_col=1, max_col=1, values_only=True)
        for c in row
        if isinstance(c, str)
    ]
    assert any("Dossier incomplet" in m for m in motifs)

    synthese = classeur["Synthèse"]
    valeurs = {
        row[0]: row[1] for row in synthese.iter_rows(min_col=1, max_col=2, values_only=True)
    }
    assert valeurs["Candidatures reçues"] == 2
    # Le vocabulaire du processus : préqualifiés d'abord, proposés ensuite.
    assert valeurs["Préqualifiés"] == 1
    assert "Proposés au client" in valeurs


async def test_l_export_porte_la_justification_du_seuil_abaisse(client, auth):
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    await client.post(
        f"{API}/postes/{poste_id}/seuil",
        json={"seuil": 12, "justification": "Poste en tension, deux candidatures."},
        headers=auth,
    )

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    classeur = load_workbook(io.BytesIO(reponse.content))
    textes = [
        c
        for row in classeur["Grille de présélection"].iter_rows(values_only=True)
        for c in row
        if isinstance(c, str)
    ]
    assert any("Poste en tension" in t for t in textes)


async def _publier_avis(client, auth, poste_id: str) -> str:
    avis = (
        await client.post(
            f"{API}/postes/{poste_id}/avis",
            json={"reference": "PUB-1", "date_cloture": "2026-12-31"},
            headers=auth,
        )
    ).json()
    publie = await client.post(f"{API}/avis/{avis['id']}/publier", headers=auth)
    assert publie.status_code == 200, publie.text
    return publie.json()["cle_publique"]


async def test_un_avis_brouillon_est_invisible_du_public(client, auth):
    poste_id = await monter_poste(client, auth)
    avis = (
        await client.post(
            f"{API}/postes/{poste_id}/avis", json={"date_cloture": "2026-12-31"}, headers=auth
        )
    ).json()

    assert (await client.get(f"{API}/public/avis")).json() == []
    # Même réponse qu'une clé inconnue : un brouillon ne se devine pas.
    assert (await client.get(f"{API}/public/avis/{avis['cle_publique']}")).status_code == 404


async def test_le_lien_d_un_avis_publie_mene_au_poste(client, auth):
    poste_id = await monter_poste(client, auth)
    cle = await _publier_avis(client, auth, poste_id)

    index = (await client.get(f"{API}/public/avis")).json()
    assert [a["cle_publique"] for a in index] == [cle]

    detail = (await client.get(f"{API}/public/avis/{cle}")).json()
    assert detail["intitule"] == "Directeur Général"
    assert detail["accepte_candidatures"] is True
    assert {p["code"] for p in detail["pieces_attendues"]} == {"LETTRE_MOTIVATION", "CV"}
    assert any("BAC+4" in ligne for ligne in detail["profil"])
    # Rien du dossier interne ne fuit vers le public.
    corps = str(detail)
    assert "seuil" not in corps and "bareme" not in corps


async def test_depot_public_cree_une_candidature_declaree(client, auth):
    poste_id = await monter_poste(client, auth)
    cle = await _publier_avis(client, auth, poste_id)

    reponse = await client.post(
        f"{API}/public/avis/{cle}/candidater",
        data={
            "nom": "Kodjo",
            "prenom": "Amina",
            "email": "amina@example.com",
            "types_pieces": ["LETTRE_MOTIVATION", "CV"],
        },
        files=[
            ("fichiers", ("lm.pdf", PDF, "application/pdf")),
            ("fichiers", ("cv.pdf", PDF, "application/pdf")),
        ],
    )
    assert reponse.status_code == 201, reponse.text
    assert reponse.json()["ok"] is True

    page = (await client.get(f"{API}/postes/{poste_id}/candidatures", headers=auth)).json()
    assert page["total"] == 1
    item = page["items"][0]
    assert item["source"] == "FORMULAIRE"
    # Le formulaire public ne demande pas de ressaisir diplômes et expériences :
    # le dossier attend d'être dépouillé, il n'est pas éliminé pour autant.
    assert item["statut"] == "A_VERIFIER"
    assert item["a_verifier"] is True


async def test_un_second_depot_avec_le_meme_email_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    cle = await _publier_avis(client, auth, poste_id)

    envoi = lambda: client.post(  # noqa: E731
        f"{API}/public/avis/{cle}/candidater",
        data={
            "nom": "Kodjo",
            "prenom": "Amina",
            "email": "amina@example.com",
            "types_pieces": ["LETTRE_MOTIVATION", "CV"],
        },
        files=[
            ("fichiers", ("lm.pdf", PDF, "application/pdf")),
            ("fichiers", ("cv.pdf", PDF, "application/pdf")),
        ],
    )
    assert (await envoi()).status_code == 201
    assert (await envoi()).status_code == 409


async def test_un_depot_sans_piece_obligatoire_est_refuse(client, auth):
    """Le contrôle porte sur les pièces exigées par l'avis, pas sur le nombre."""
    poste_id = await monter_poste(client, auth)
    cle = await _publier_avis(client, auth, poste_id)

    reponse = await client.post(
        f"{API}/public/avis/{cle}/candidater",
        data={
            "nom": "Kodjo",
            "prenom": "Amina",
            "email": "amina@example.com",
            "types_pieces": ["CV"],  # lettre de motivation manquante
        },
        files=[("fichiers", ("cv.pdf", PDF, "application/pdf"))],
    )
    assert reponse.status_code == 422
    assert "Lettre de motivation" in reponse.json()["detail"]


async def test_une_piece_hors_nomenclature_est_refusee(client, auth):
    poste_id = await monter_poste(client, auth)
    cle = await _publier_avis(client, auth, poste_id)

    reponse = await client.post(
        f"{API}/public/avis/{cle}/candidater",
        data={
            "nom": "Kodjo",
            "prenom": "Amina",
            "email": "amina@example.com",
            "types_pieces": ["LETTRE_MOTIVATION", "CV", "CASIER_JUDICIAIRE"],
        },
        files=[
            ("fichiers", ("lm.pdf", PDF, "application/pdf")),
            ("fichiers", ("cv.pdf", PDF, "application/pdf")),
            ("fichiers", ("casier.pdf", PDF, "application/pdf")),
        ],
    )
    assert reponse.status_code == 422
    assert "non prévue" in reponse.json()["detail"]


async def test_un_avis_cloture_refuse_les_depots(client, auth):
    poste_id = await monter_poste(client, auth)
    avis = (
        await client.post(
            f"{API}/postes/{poste_id}/avis",
            # Clôture dépassée : le dépôt se ferme sans intervention.
            json={"date_cloture": "2020-01-01"},
            headers=auth,
        )
    ).json()
    await client.post(f"{API}/avis/{avis['id']}/publier", headers=auth)
    cle = avis["cle_publique"]

    assert (await client.get(f"{API}/public/avis")).json() == []
    detail = (await client.get(f"{API}/public/avis/{cle}")).json()
    assert detail["accepte_candidatures"] is False

    refus = await client.post(
        f"{API}/public/avis/{cle}/candidater",
        data={"nom": "X", "prenom": "Y", "email": "x@y.com", "types_pieces": ["CV"]},
        files=[("fichiers", ("cv.pdf", PDF, "application/pdf"))],
    )
    assert refus.status_code == 409


async def test_authentification_requise(client):
    reponse = await client.get(f"{API}/clients")
    assert reponse.status_code == 401


# --- la fiche de poste se modifie après coup ---------------------------------


async def test_modifier_une_exigence_reevalue_le_poste(client, auth):
    """La grille est ce que le client relit : elle ne peut pas rester en arrière.

    Sans cette réévaluation, changer le niveau exigé laissait les motifs
    d'élimination d'avant la modification, et il fallait penser à cliquer
    « Recalculer » pour que la grille corresponde à la fiche.
    """
    poste_id = await monter_poste(client, auth)

    reponse = await client.patch(
        f"{API}/postes/{poste_id}", json={"niveau_min": 5}, headers=auth
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["niveau_min"] == 5

    async with SessionLocal() as db:
        journal = (
            await db.execute(select(AuditLog).where(AuditLog.action == "poste.update"))
        ).scalars().all()
    assert journal, "une modification de fiche se journalise"
    assert "candidatures_reevaluees" in journal[-1].details


async def test_plusieurs_experiences_specifiques_se_declarent(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.patch(
        f"{API}/postes/{poste_id}",
        json={
            "annees_experience_specifique_min": 0,
            "domaines_experience": [],
            "experiences_specifiques": [
                {"libelle": "passation des marchés", "domaines": ["marches"], "annees_min": 5},
                {"libelle": "gestion de projet", "domaines": ["projet"], "annees_min": 3,
                 "poids": 2},
            ],
        },
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    exigences = reponse.json()["experiences_specifiques"]
    assert [e["libelle"] for e in exigences] == ["passation des marchés", "gestion de projet"]
    assert exigences[1]["poids"] == 2


async def test_une_exigence_specifique_ne_depasse_pas_l_experience_generale(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.patch(
        f"{API}/postes/{poste_id}",
        json={
            "annees_experience_min": 5,
            "experiences_specifiques": [
                {"libelle": "trop long", "domaines": ["x"], "annees_min": 12}
            ],
        },
        headers=auth,
    )
    assert reponse.status_code == 422


async def test_une_exigence_doit_pouvoir_se_nommer(client, auth):
    """Sans libellé ni domaine, ni la grille ni le candidat ne savent de quoi
    il s'agit."""
    poste_id = await monter_poste(client, auth)
    reponse = await client.patch(
        f"{API}/postes/{poste_id}",
        json={"experiences_specifiques": [{"annees_min": 3}]},
        headers=auth,
    )
    assert reponse.status_code == 422


async def test_les_extras_de_formation_restent_dans_les_sept_points(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.put(
        f"{API}/postes/{poste_id}/bareme/formation",
        json={
            "points_par_certification": 0.5,
            "certifications_max": 3,
            "points_formation_complementaire": 1.0,
        },
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    formation = reponse.json()["bareme"]["formation"]
    assert formation["points_par_certification"] == 0.5
    assert formation["points_formation_complementaire"] == 1.0
    # Le maximum de la ligne n'a pas bougé : le barème vaut toujours 30.
    assert formation["points_max"] == 7.0
    assert reponse.json()["bareme"]["total_max"] == 30.0


async def test_les_destinataires_se_listent_par_portee(client, auth):
    poste_id = await monter_poste(client, auth)
    reponse = await client.get(f"{API}/postes/{poste_id}/destinataires", headers=auth)
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["portee"] == "tous"
    assert reponse.json()["candidature_ids"] == []

    refus = await client.get(
        f"{API}/postes/{poste_id}/destinataires?portee=inconnue", headers=auth
    )
    assert refus.status_code == 422


async def test_deplacer_le_seuil_reevalue_la_preselection(client, auth):
    """C'est le seul objet de ce geste.

    L'écran montre la distribution des notes pour aider à placer la barre ;
    la barre bougeait, et la liste des préqualifiés restait celle de l'ancienne
    jusqu'à ce que quelqu'un pense à cliquer « Recalculer ».
    """
    poste_id = await monter_poste(client, auth)

    reponse = await client.post(
        f"{API}/postes/{poste_id}/seuil",
        json={"seuil": 12, "justification": "Barre posée au vu des notes reçues."},
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["seuil_preselection"] == 12

    async with SessionLocal() as db:
        journal = (
            await db.execute(select(AuditLog).where(AuditLog.action == "poste.seuil"))
        ).scalars().all()
    assert journal, "un déplacement de barre se journalise"
    assert "candidatures_reevaluees" in journal[-1].details


async def test_abaisser_le_seuil_sans_justification_est_refuse(client, auth):
    """Abaisser la barre change qui est écarté : cela ne peut pas être implicite."""
    poste_id = await monter_poste(client, auth)
    await client.post(
        f"{API}/postes/{poste_id}/seuil",
        json={"seuil": 20, "justification": "Première barre, posée sur la distribution."},
        headers=auth,
    )
    refus = await client.post(
        f"{API}/postes/{poste_id}/seuil", json={"seuil": 10, "justification": ""}, headers=auth
    )
    assert refus.status_code == 422
