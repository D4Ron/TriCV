"""Entretiens structurés : la seconde étape, et la note sur 100.

Le point que ces tests tiennent : ces 70 points sont un jugement humain. Rien
ne les calcule, rien ne les propose, et une saisie partielle ne se lit jamais
comme un résultat définitif.
"""

from __future__ import annotations

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.domain import BAREME_ENTRETIEN
from tests.test_api_recrutement import dossier, monter_poste

pytestmark = pytest.mark.anyio

API = "/api/v1"

# Une séance complète : chaque critère noté au maximum.
PARFAIT = [{"code": l.code, "points": l.points_max} for l in BAREME_ENTRETIEN]


async def deposer(client, auth, poste_id: str, **overrides) -> dict:
    reponse = await client.post(
        f"{API}/postes/{poste_id}/candidatures", json=dossier(**overrides), headers=auth
    )
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


# --- la fiche ----------------------------------------------------------------


async def test_la_fiche_presente_la_grille_entiere_avant_toute_saisie(client, auth):
    """Un jury doit voir ce qu'il lui reste à noter, pas le deviner."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    fiche = (
        await client.get(f"{API}/candidatures/{corps['id']}/entretien", headers=auth)
    ).json()

    assert fiche["existe"] is False
    assert fiche["total_max"] == 70.0
    assert [l["code"] for l in fiche["grille"]] == [l.code for l in BAREME_ENTRETIEN]
    assert fiche["fiches"] == []
    assert fiche["complet"] is False


async def test_la_grille_reprend_la_repartition_du_cabinet(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    fiche = (
        await client.get(f"{API}/candidatures/{corps['id']}/entretien", headers=auth)
    ).json()
    maxima = {l["code"]: l["points_max"] for l in fiche["grille"]}
    # Relevé dans l'offre technique du cabinet, et non déduit.
    assert maxima == {
        "PRESENTATION": 2.0,
        "MOTIVATION": 2.0,
        "RELATIONNELLES": 20.0,
        "TECHNIQUES": 25.0,
        "POTENTIEL": 20.0,
        "CONNAISSANCES_CLIENT": 1.0,
    }
    assert sum(maxima.values()) == 70.0


# --- saisie ------------------------------------------------------------------


async def test_un_entretien_complet_donne_la_note_sur_cent(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    preselection = corps["notation"]["total"]

    reponse = await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={
            "lignes": PARFAIT,
            "date_entretien": "2026-09-15",
            "jury": "M. Adjovi, Mme Lawson",
            "observations": "Prestation solide.",
        },
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    fiche = reponse.json()

    assert fiche["existe"] is True
    assert fiche["total"] == 70.0
    assert fiche["complet"] is True
    assert fiche["entretien_sur_cent"] == 70.0
    # 30 % pour la présélection, 70 pour l'entretien.
    assert fiche["preselection_sur_cent"] == pytest.approx(preselection, abs=0.5)
    assert fiche["note_finale_sur_cent"] == pytest.approx(
        fiche["preselection_sur_cent"] + 70.0
    )


async def test_une_saisie_partielle_reste_annoncee_partielle(client, auth):
    """Un acquis en cours de séance ne doit pas se lire comme un résultat."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"lignes": [{"code": "TECHNIQUES", "points": 24}]},
            headers=auth,
        )
    ).json()

    assert fiche["complet"] is False
    assert fiche["total"] == 24.0
    assert 0 < fiche["entretien_sur_cent"] < 70.0


async def test_le_commentaire_accompagne_la_note(client, auth):
    """Une note se défend par ce qui a été observé, pas par le chiffre."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={
                "lignes": [
                    {
                        "code": "POTENTIEL",
                        "points": 9,
                        "commentaire": "Exemples concrets de conduite d'équipe.",
                    }
                ]
            },
            headers=auth,
        )
    ).json()

    lignes = fiche["fiches"][0]["lignes"]
    ligne = next(l for l in lignes if l["code"] == "POTENTIEL")
    assert ligne["points"] == 9.0
    assert "Exemples concrets" in ligne["commentaire"]


async def test_la_fiche_est_un_etat_et_non_un_journal(client, auth):
    """Réenvoyer la fiche remplace les lignes : une note retirée disparaît."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )
    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"lignes": [{"code": "TECHNIQUES", "points": 10}]},
            headers=auth,
        )
    ).json()

    assert fiche["total"] == 10.0
    assert fiche["complet"] is False
    notes = {l["code"]: l["points"] for l in fiche["fiches"][0]["lignes"]}
    assert notes["TECHNIQUES"] == 10.0
    assert notes["POTENTIEL"] is None


# --- refus -------------------------------------------------------------------


async def test_une_note_hors_bareme_est_refusee(client, auth):
    """Borner en silence ferait afficher une note que le jury n'a pas mise."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    refus = await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": [{"code": "CONNAISSANCES_CLIENT", "points": 12}]},
        headers=auth,
    )
    assert refus.status_code == 409
    assert "hors barème" in refus.json()["detail"]


async def test_un_critere_inconnu_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    refus = await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": [{"code": "SYMPATHIE", "points": 5}]},
        headers=auth,
    )
    assert refus.status_code == 409
    assert "inconnu" in refus.json()["detail"]


async def test_un_dossier_elimine_ne_se_note_pas(client, auth):
    """Convoquer un candidat écarté se décide ; cela ne se contourne pas."""
    poste_id = await monter_poste(client, auth)
    incomplet = await deposer(client, auth, poste_id, pieces_fournies=["CV"])
    assert incomplet["statut"] == "ELIMINEE"

    refus = await client.put(
        f"{API}/candidatures/{incomplet['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )
    assert refus.status_code == 409
    assert "Levez le motif" in refus.json()["detail"]


async def test_l_entretien_devient_possible_une_fois_le_motif_leve(client, auth):
    poste_id = await monter_poste(client, auth)
    incomplet = await deposer(client, auth, poste_id, pieces_fournies=["CV"])
    motif = incomplet["eliminations"][0]["motif"]

    await client.post(
        f"{API}/candidatures/{incomplet['id']}/eliminations/{motif}/lever",
        json={"motif": "Pièce reçue par courrier séparé."},
        headers=auth,
    )

    reponse = await client.put(
        f"{API}/candidatures/{incomplet['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text


# --- effacement et journal ---------------------------------------------------


async def test_une_fiche_saisie_par_erreur_s_efface(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )

    reponse = await client.request(
        "DELETE", f"{API}/candidatures/{corps['id']}/entretien", headers=auth
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["existe"] is False
    assert reponse.json()["total"] == 0.0


async def test_effacer_une_fiche_absente_renvoie_404(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    reponse = await client.request(
        "DELETE", f"{API}/candidatures/{corps['id']}/entretien", headers=auth
    )
    assert reponse.status_code == 404


async def test_la_saisie_et_l_effacement_sont_journalises(client, auth):
    from app.models import AuditLog

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT, "jury": "Comité de sélection"},
        headers=auth,
    )
    await client.request(
        "DELETE", f"{API}/candidatures/{corps['id']}/entretien", headers=auth
    )

    async with SessionLocal() as db:
        actions = [a.action for a in (await db.execute(select(AuditLog))).scalars()]
    assert "candidature.entretien" in actions
    assert "candidature.entretien_efface" in actions


# --- la grille porte la note finale ------------------------------------------


async def test_la_grille_porte_la_note_finale(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )

    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    assert grille["nombre_entretiens"] == 1

    ligne = grille["preselectionnes"][0]
    assert ligne["note_entretien_sur_cent"] == 70.0
    assert ligne["entretien_complet"] is True
    assert ligne["note_finale_sur_cent"] == pytest.approx(
        ligne["note_sur_cent"] + 70.0
    )


async def test_sans_entretien_la_grille_n_invente_pas_de_zero(client, auth):
    """« Pas encore reçu » et « n'a rien obtenu » ne se confondent pas."""
    poste_id = await monter_poste(client, auth)
    await deposer(client, auth, poste_id)

    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    ligne = grille["preselectionnes"][0]
    assert ligne["note_entretien_sur_cent"] is None
    assert ligne["note_finale_sur_cent"] is None
    assert grille["nombre_entretiens"] == 0


async def test_l_export_porte_la_note_finale(client, auth):
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    assert reponse.status_code == 200, reponse.text
    classeur = load_workbook(io.BytesIO(reponse.content))

    grille = classeur["Grille de présélection"]
    entetes = {
        c for row in grille.iter_rows(values_only=True) for c in row if isinstance(c, str)
    }
    assert "Note finale /100" in entetes

    synthese = classeur["Synthèse"]
    valeurs = {
        row[0]: row[1] for row in synthese.iter_rows(min_col=1, max_col=2, values_only=True)
    }
    assert valeurs["Entretiens saisis"] == 1


async def test_un_entretien_partiel_est_signale_dans_l_export(client, auth):
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": [{"code": "TECHNIQUES", "points": 20}]},
        headers=auth,
    )

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    classeur = load_workbook(io.BytesIO(reponse.content))
    grille = classeur["Grille de présélection"]
    contenu = [
        c for row in grille.iter_rows(values_only=True) for c in row if isinstance(c, str)
    ]
    assert any("(partiel)" in c for c in contenu)


# --- la grille d'entretien est figée -----------------------------------------


async def test_la_grille_est_figee_a_la_saisie(client, auth):
    """Changer la répartition des 70 points ne réécrit pas une fiche rendue."""
    from app.models import Entretien

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": PARFAIT},
        headers=auth,
    )

    async with SessionLocal() as db:
        entretien = (await db.execute(select(Entretien))).scalar_one()
        fige = {l["code"]: l["points_max"] for l in entretien.bareme_utilise}

    assert fige == {l.code: l.points_max for l in BAREME_ENTRETIEN}


# --- matière pour le rapport -------------------------------------------------


async def test_l_export_porte_le_detail_des_entretiens(client, auth):
    """L'application ne rédige pas le rapport ; elle en fournit la matière."""
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    saisie = await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={
            "lignes": [
                {
                    "code": "TECHNIQUES",
                    "points": 24,
                    "commentaire": "Maîtrise du contrôle de gestion hospitalier.",
                },
                {"code": "RELATIONNELLES", "points": 15},
                {"code": "POTENTIEL", "points": 9},
                {"code": "MOTIVATION", "points": 2},
                {"code": "PRESENTATION", "points": 2},
                {"code": "CONNAISSANCES_CLIENT", "points": 1},
            ],
            "date_entretien": "2026-09-15",
            "jury": "M. Adjovi, Mme Lawson",
            "observations": "Vision claire du poste.",
        },
        headers=auth,
    )
    assert saisie.status_code == 200, saisie.text

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    classeur = load_workbook(io.BytesIO(reponse.content))
    assert "Entretiens" in classeur.sheetnames

    feuille = classeur["Entretiens"]
    contenu = [
        c for row in feuille.iter_rows(values_only=True) for c in row if isinstance(c, str)
    ]
    # Ce qui sert à écrire le rapport : le jury, la date, et ce qui a été observé.
    assert any("M. Adjovi" in c for c in contenu)
    assert any("contrôle de gestion hospitalier" in c for c in contenu)
    assert any("Vision claire du poste" in c for c in contenu)
    assert any("Total entretien" in c for c in contenu)


async def test_l_onglet_entretiens_n_apparait_pas_sans_entretien(client, auth):
    """Un onglet vide dans un fichier remis au client interroge pour rien."""
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    await deposer(client, auth, poste_id)

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    classeur = load_workbook(io.BytesIO(reponse.content))
    assert "Entretiens" not in classeur.sheetnames


async def test_un_critere_non_note_est_dit_tel_quel_dans_l_export(client, auth):
    """Une case vide laisserait croire à un oubli de lecture, pas de notation."""
    import io

    from openpyxl import load_workbook

    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)
    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"lignes": [{"code": "TECHNIQUES", "points": 20}]},
        headers=auth,
    )

    reponse = await client.get(f"{API}/postes/{poste_id}/grille.xlsx", headers=auth)
    classeur = load_workbook(io.BytesIO(reponse.content))
    contenu = [
        c
        for row in classeur["Entretiens"].iter_rows(values_only=True)
        for c in row
        if isinstance(c, str)
    ]
    assert any("Non noté" in c for c in contenu)
    assert any("partiel" in c for c in contenu)


# --- panel de jurés ----------------------------------------------------------


async def test_plusieurs_jures_donnent_une_moyenne(client, auth):
    """Le rapport du cabinet publie une « Moyenne/100 » d'un panel de sept.

    Une fiche unique par candidature aurait écrasé six avis sur sept.
    """
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    for jure, technique in (("Mme Adjovi", 20), ("M. Lawson", 25), ("M. Panassi", 15)):
        reponse = await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"jure": jure, "lignes": [{"code": "TECHNIQUES", "points": technique}]},
            headers=auth,
        )
        assert reponse.status_code == 200, reponse.text

    fiche = reponse.json()
    assert len(fiche["fiches"]) == 3
    assert [f["jure"] for f in fiche["fiches"]] == ["M. Lawson", "M. Panassi", "Mme Adjovi"]
    # Moyenne de 20, 25 et 15.
    assert fiche["total"] == 20.0
    # L'écart signale un panel dispersé, à relire avant que la note ne parte.
    assert fiche["ecart_jures"] == 10.0


async def test_la_fiche_d_un_jure_ne_touche_pas_celle_des_autres(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"jure": "Premier", "lignes": PARFAIT},
        headers=auth,
    )
    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"jure": "Second", "lignes": [{"code": "TECHNIQUES", "points": 5}]},
            headers=auth,
        )
    ).json()

    par_jure = {f["jure"]: f["total"] for f in fiche["fiches"]}
    assert par_jure == {"Premier": 70.0, "Second": 5.0}


async def test_effacer_la_fiche_d_un_seul_jure(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    for jure in ("Premier", "Second"):
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"jure": jure, "lignes": PARFAIT},
            headers=auth,
        )

    reponse = await client.request(
        "DELETE",
        f"{API}/candidatures/{corps['id']}/entretien?jure=Premier",
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    assert [f["jure"] for f in reponse.json()["fiches"]] == ["Second"]


async def test_le_panel_n_est_complet_que_si_chaque_fiche_l_est(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    await client.put(
        f"{API}/candidatures/{corps['id']}/entretien",
        json={"jure": "Complet", "lignes": PARFAIT},
        headers=auth,
    )
    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"jure": "Partiel", "lignes": [{"code": "TECHNIQUES", "points": 10}]},
            headers=auth,
        )
    ).json()

    assert fiche["complet"] is False


# --- grille négociée par poste ------------------------------------------------


async def test_la_grille_se_regle_poste_par_poste(client, auth):
    """L'offre technique la dit « indicative », à valider par le client.

    Un mandat réel s'en écarte : celui de 2025 notait sur quatre rubriques
    totalisant 100, avec des critères propres au poste.
    """
    poste_id = await monter_poste(client, auth)

    grille = {
        "criteres": [
            {"code": "PRESENTATION", "libelle": "Présentation", "points_max": 3,
             "section": "I. Présentation et motivations"},
            {"code": "MOTIVATION", "libelle": "Motivation", "points_max": 2,
             "section": "I. Présentation et motivations"},
            {"code": "BUDGET", "libelle": "Gestion budgétaire", "points_max": 55,
             "section": "II. Expérience et compétences"},
            {"code": "MANAGEMENT", "libelle": "Aptitude managériale", "points_max": 40,
             "section": "III. Capacités managériales"},
        ]
    }
    reponse = await client.put(
        f"{API}/postes/{poste_id}/grille-entretien", json=grille, headers=auth
    )
    assert reponse.status_code == 200, reponse.text

    corps = await deposer(client, auth, poste_id)
    fiche = (
        await client.get(f"{API}/candidatures/{corps['id']}/entretien", headers=auth)
    ).json()

    assert [l["code"] for l in fiche["grille"]] == [
        "PRESENTATION", "MOTIVATION", "BUDGET", "MANAGEMENT"
    ]
    assert fiche["total_max"] == 100.0
    # Les rubriques et leur total, comme dans le document du cabinet.
    assert fiche["sections"] == [
        {"libelle": "I. Présentation et motivations", "points_max": 5.0},
        {"libelle": "II. Expérience et compétences", "points_max": 55.0},
        {"libelle": "III. Capacités managériales", "points_max": 40.0},
    ]


async def test_un_critere_propre_au_poste_se_note(client, auth):
    poste_id = await monter_poste(client, auth)
    await client.put(
        f"{API}/postes/{poste_id}/grille-entretien",
        json={"criteres": [
            {"code": "BUDGET", "libelle": "Gestion budgétaire", "points_max": 70}
        ]},
        headers=auth,
    )
    corps = await deposer(client, auth, poste_id)

    fiche = (
        await client.put(
            f"{API}/candidatures/{corps['id']}/entretien",
            json={"lignes": [{"code": "BUDGET", "points": 60}]},
            headers=auth,
        )
    ).json()
    assert fiche["total"] == 60.0
    assert fiche["complet"] is True


async def test_une_grille_a_codes_doubles_est_refusee(client, auth):
    poste_id = await monter_poste(client, auth)
    refus = await client.put(
        f"{API}/postes/{poste_id}/grille-entretien",
        json={"criteres": [
            {"code": "X", "libelle": "Un", "points_max": 30},
            {"code": "X", "libelle": "Deux", "points_max": 40},
        ]},
        headers=auth,
    )
    assert refus.status_code == 422
    assert "même code" in refus.json()["detail"]


async def test_remettre_la_grille_du_cabinet(client, auth):
    poste_id = await monter_poste(client, auth)
    await client.put(
        f"{API}/postes/{poste_id}/grille-entretien",
        json={"criteres": [{"code": "X", "libelle": "Un", "points_max": 70}]},
        headers=auth,
    )
    reponse = await client.put(
        f"{API}/postes/{poste_id}/grille-entretien", json={"criteres": None}, headers=auth
    )
    assert reponse.status_code == 200
    assert [c["code"] for c in reponse.json()["criteres"]] == [
        l.code for l in BAREME_ENTRETIEN
    ]


# --- qualification remise au client -------------------------------------------


async def test_un_dossier_note_recoit_une_categorie(client, auth):
    """Fortement / partiellement / non qualifié : ce que le client reçoit."""
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    assert corps["qualification"] in {
        "FORTEMENT_QUALIFIE", "PARTIELLEMENT_QUALIFIE", "NON_QUALIFIE"
    }
    assert corps["qualification_libelle"]


async def test_la_categorie_se_reinscrit_avec_un_motif(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    reponse = await client.patch(
        f"{API}/candidatures/{corps['id']}/qualification",
        json={"qualification": "NON_QUALIFIE", "motif": "Références défavorables."},
        headers=auth,
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["qualification"] == "NON_QUALIFIE"
    assert reponse.json()["qualification_libelle"] == "Non qualifié"


async def test_reinscrire_sans_motif_est_refuse(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    refus = await client.patch(
        f"{API}/candidatures/{corps['id']}/qualification",
        json={"qualification": "NON_QUALIFIE", "motif": ""},
        headers=auth,
    )
    assert refus.status_code == 422


async def test_une_categorie_inconnue_est_refusee(client, auth):
    poste_id = await monter_poste(client, auth)
    corps = await deposer(client, auth, poste_id)

    refus = await client.patch(
        f"{API}/candidatures/{corps['id']}/qualification",
        json={"qualification": "EXCELLENT", "motif": "Essai."},
        headers=auth,
    )
    assert refus.status_code == 422
    assert "inconnue" in refus.json()["detail"]


async def test_la_grille_porte_la_categorie(client, auth):
    poste_id = await monter_poste(client, auth)
    await deposer(client, auth, poste_id)

    grille = (await client.get(f"{API}/postes/{poste_id}/grille", headers=auth)).json()
    ligne = grille["preselectionnes"][0]
    assert ligne["qualification"]
    assert ligne["qualification_libelle"]
